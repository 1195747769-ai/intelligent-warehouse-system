"""Local document intake. Source evidence is retained; corrections are replayed on import."""
import base64
import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl


LABELS = dict(code='厂家编码', name='物资名称', spec='规格型号', drawing='图号', unit='单位', qty='数量',
              box='箱号', doc_no='到货单号', attribute='备注')
EDITABLE = set(LABELS) | {'exclude_reason'}
MAX_ROWS = 10000
BOX_PATTERN = r'(?:发\s*运\s*箱\s*号|包\s*装\s*箱\s*号|箱\s*件\s*号|箱\s*号)\s*[:：]\s*([A-Za-z0-9][\w./-]*)'


def packed(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')


def text_rows(text):
    if not isinstance(text, str) or not text.strip() or len(text.encode('utf-8')) > 12 * 1024 * 1024:
        raise ValueError('内容为空或超过12MB，请分批录入')
    if '\t' in text:
        return list(csv.reader(io.StringIO(text), delimiter='\t'))
    if ',' in text or '，' in text:
        return list(csv.reader(io.StringIO(text.replace('，', ','))))
    # ponytail: whitespace PDF extraction is best effort; ambiguous lines stay editable.
    return [re.split(r'\s{2,}|\s*\|\s*', line.strip()) for line in text.splitlines()]


def inferred_spec(value):
    patterns = (
        r'(?<![A-Za-z0-9])M\s*\d+(?:\.\d+)?(?:\s*[xX×]\s*\d+(?:\.\d+)?)*(?:-[A-Za-z0-9]+)*',
        r'(?<![A-Za-z0-9])DN\s*\d+', r'(?<![A-Za-z0-9])G\s*\d+(?:/\d+)?',
        r'(?<![A-Za-z0-9])[φΦØø]\s*\d+(?:\.\d+)?',
        r'(?<![A-Za-z0-9])\d+(?:\.\d+)?\s*[xX×]\s*\d+(?:\.\d+)?(?:\s*[xX×]\s*\d+(?:\.\d+)?)*(?:-[A-Za-z0-9]+)*',
        r'(?<![A-Za-z0-9])\d{1,3}[A-Z]{1,2}\d{1,3}(?:[A-Z]{1,2}\d{0,3}){1,3}(?![A-Za-z0-9])',
        r'(?<![A-Za-z0-9])\d+(?:\.\d+)?\s*(?:t|吨)(?![A-Za-z0-9])',
        r'(?<![A-Za-z0-9])(?:AWS\s+)?[A-Z]{1,5}\d{1,5}(?:\.\d+)?[A-Za-z0-9]*(?:-[A-Za-z0-9]+)*(?![A-Za-z0-9])',
    )
    matches=[]
    for pattern in patterns:
        matches.extend((m.start(),m.end(),m.group().strip()) for m in re.finditer(pattern,value or '',re.I))
    selected=[]
    for start,end,text in sorted(matches,key=lambda match:(match[0],-match[1])):
        if any(start<b and end>a for a,b,_ in selected): continue
        selected.append((start,end,text))
    return ' '.join(text for _,_,text in selected)


def read_input(source, app):
    mode = source.get('source')
    if mode == 'manual':
        records = source.get('rows')
        if not isinstance(records, list) or not all(isinstance(r, dict) for r in records):
            raise ValueError('手动明细格式无效')
        return {'手动录入': {'rows': [list(LABELS.values())] + [[app.clean(r.get(k)) for k in LABELS] for r in records], 'merged': []}}
    if mode == 'paste':
        return {'粘贴内容': {'rows': text_rows(source.get('text')), 'merged': []}}
    if mode != 'upload':
        raise ValueError('请选择上传、粘贴或手动录入')
    filename = app.required(source.get('filename'), '文件名')
    if Path(filename).name.startswith('~$'):
        raise ValueError('这是 Excel 锁文件，请上传不以 ~$ 开头的实际工作簿')
    try:
        raw = base64.b64decode(source.get('content', ''), validate=True)
    except Exception:
        raise ValueError('文件编码无效') from None
    if not raw or len(raw) > 12 * 1024 * 1024:
        raise ValueError('文件为空或超过12MB')
    ext = Path(filename).suffix.lower()
    sheets = {}
    try:
        if ext == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(x.file_size for x in z.infolist()) > 80 * 1024 * 1024:
                    raise ValueError('解压文件过大')
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
            try:
                for ws in wb:
                    if ws.max_row > MAX_ROWS + 100 or ws.max_column > 200:
                        raise ValueError('工作表超过10000行或200列，请拆分后上传')
                    rows = [[{'formula': True} if c.data_type == 'f' else app.clean(c.value) for c in row] for row in ws]
                    merged = [(r.min_row, r.max_row, r.min_col, r.max_col) for r in ws.merged_cells.ranges]
                    sheets[ws.title] = dict(rows=rows, merged=merged)
            finally:
                wb.close()
        elif ext == '.xls':
            import xlrd
            wb = xlrd.open_workbook(file_contents=raw, formatting_info=True)
            try:
                for ws in wb.sheets():
                    if ws.nrows > MAX_ROWS + 100 or ws.ncols > 200:
                        raise ValueError('工作表超过10000行或200列，请拆分后上传')
                    rows = [[{'formula': True} if c.ctype == xlrd.XL_CELL_ERROR else app.clean(int(c.value) if c.ctype == xlrd.XL_CELL_NUMBER and c.value.is_integer() else c.value) for c in ws.row(i)] for i in range(ws.nrows)]
                    sheets[ws.name] = dict(rows=rows, merged=[(a+1, b, c+1, d) for a,b,c,d in ws.merged_cells])
            finally:
                wb.release_resources()
        elif ext == '.csv':
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = raw.decode('gb18030')
            sheets['CSV'] = dict(rows=text_rows(text), merged=[])
        else:
            raise ValueError('支持 .xls、.xlsx、.csv；其他内容可复制到粘贴入口')
    except (ValueError, ImportError):
        raise
    except Exception:
        raise ValueError('文件无法读取，请检查文件是否损坏、加密或格式与扩展名不符') from None
    return sheets


def header_info(rows, app, requested=0):
    aliases = {k: {app.norm(a) for a in vals} for k, vals in app.ALIASES.items()}
    aliases['doc_no'] = {app.norm(a) for a in ('到货单号', '送货单号', '收货单号', '到货单编号', 'deliveryno', 'deliverynotenumber')}
    aliases['name'].update(map(app.norm, ('产品名称', '货品名称', '材料品名', '零部件名称')))
    aliases['qty'].update(map(app.norm, ('件数', '本箱数量', '装箱数量', '实到数量', '发货数量')))
    aliases['spec'] -= {app.norm(x) for x in ('图号','零件图号','图纸编号','drawingno','drawingnumber')}
    aliases['drawing'] = {app.norm(x) for x in ('图号','零件图号','图纸编号','drawingno','drawingnumber')}
    def infer(row):
        mapping = {}
        for key in LABELS:
            hits = [i for i, cell in enumerate(row) if not isinstance(cell, dict) and app.norm(cell) in aliases.get(key, set())]
            mapping[key] = hits[0] if len(hits) == 1 or (key == 'attribute' and hits) else -1
        return mapping
    limit = min(len(rows), 100)
    options = []
    for i in range(limit):
        for depth in (1, 2):
            if i+depth > len(rows): continue
            headers = [app.clean(x) if not isinstance(x, dict) else '' for x in rows[i]]
            if depth == 2:
                next_row = rows[i+1]
                headers = [''.join(dict.fromkeys([headers[c] if c<len(headers) else '', app.clean(next_row[c]) if c<len(next_row) else ''])) for c in range(max(len(headers),len(next_row)))]
            mapping = infer(headers)
            score = sum(mapping[k] >= 0 for k in ('name','qty','unit','box','doc_no'))
            if mapping['name'] >= 0: score += 3
            options.append((score, i, depth, headers, mapping))
    if not options:
        return dict(header_row=1, end=1, headers=[], mapping={k:-1 for k in LABELS}, score=0)
    if requested:
        options = [o for o in options if o[1] == requested-1]
        if not options: raise ValueError('表头行无效（可选前100行）')
    best = max(options, key=lambda o: (o[0], -o[2], -o[1]))
    score, i, depth, headers, mapping = best
    return dict(header_row=i+1, end=i+depth, headers=headers, mapping=mapping, score=score)


def classify(sheet, rows, info, app):
    heading = ' '.join(app.clean(c) for row in rows[:max(8, info['end'])] for c in row if not isinstance(c, dict)).lower()
    title = sheet.lower()
    headers = {app.norm(h) for h in info['headers']}
    if any(word in title for word in ('出厂资料','技术资料','文件清单','document')):
        return 'document_only', '工作表为资料/证书清单'
    packing = any(word in title+heading for word in ('装箱清单','装箱单','箱单','packing'))
    logistics = any(app.norm(w) in headers for w in ('毛重','净重','箱数','箱体尺寸','包装类型'))
    if logistics or (not packing and any(w in title+heading[:300] for w in ('发运单','发货清单','发运清单','shipment','shipping list'))):
        return 'shipment', '表名/表头包含发运或箱级物流字段'
    arrival_title = any(word in title for word in ('到货单','送货单','arrival note','delivery note'))
    arrival_heading = any(app.norm(cell) in {app.norm(word) for word in ('到货单','到货清单','送货单','送货清单','arrival note','delivery note')}
                          for row in rows[:8] for cell in row if not isinstance(cell,dict))
    if (arrival_title or arrival_heading) and not packing:
        return 'arrival_note', '表名/标题指向到货或送货单'
    if any(word in title for word in ('总清单','汇总','总表')):
        return 'material_summary', '工作表为汇总层级，须确认入库范围'
    if packing or all(info['mapping'][k] >= 0 for k in ('name','qty','unit')):
        return 'packing_detail', '装箱标题或名称/数量/单位字段已识别'
    return 'unknown', '字段不完整，请映射并确认物资层级'


def parse(p, app):
    inputs = p.get('inputs')
    if not isinstance(inputs, list) or not 1 <= len(inputs) <= 4 or not all(isinstance(s, dict) for s in inputs):
        raise ValueError('请选择1至4个录入来源')
    corrections = p.get('corrections') or {}
    if not isinstance(corrections, dict): raise ValueError('人工确认内容格式无效')
    errors, warnings, source_rows, candidates, selections, layout_rows = [], [], [], [], [], []
    originals = [{k: s[k] for k in ('source','role','filename','content','text','rows') if k in s} for s in inputs]
    raw = packed({'format': 'warehouse-intake-v1', 'sources': originals})
    if len(raw) > 17 * 1024 * 1024: raise ValueError('本次来源总大小过大，请分批录入')
    if b'DEMO-ONLY' in raw and not app.demo_mode():
        raise ValueError('演示数据只能在演示模式中使用')
    for index, source in enumerate(inputs):
        role = source.get('role')
        if role not in ('auto', 'arrival_note', 'packing_detail'): raise ValueError('请选择到货单或装箱单')
        sheets = read_input(source, app)
        if any(len(s['rows']) > MAX_ROWS+100 for s in sheets.values()): raise ValueError('最多10000条明细')
        infos = {}
        for name, data in sheets.items():
            info = header_info(data['rows'], app)
            kind, reason = classify(name, data['rows'], info, app)
            infos[name] = info
            candidates.append(dict(input=index, sheet=name, kind=kind, reason=reason, header_row=info['header_row'], mapping=info['mapping'], count=max(0,len(data['rows'])-info['end'])))
        shipment_boxes = set()
        for candidate in candidates:
            if candidate['input'] != index or candidate['kind'] != 'shipment': continue
            box_col = candidate['mapping']['box']
            if box_col >= 0:
                for cells in sheets[candidate['sheet']]['rows'][infos[candidate['sheet']]['end']:]:
                    qty_col,name_col=candidate['mapping']['qty'],candidate['mapping']['name']
                    if not (0<=qty_col<len(cells) and 0<=name_col<len(cells)) or not cells[name_col]: continue
                    try: app.units(cells[qty_col])
                    except (ValueError,TypeError): continue
                    if box_col < len(cells) and not isinstance(cells[box_col],dict):
                        box = app.clean(cells[box_col])
                        if re.fullmatch(r'[A-Za-z0-9][\w./-]*',box): shipment_boxes.add(box)
        selected = source.get('sheet') or next((c['sheet'] for c in candidates if c['input']==index and (c['kind']==role or role=='auto' and c['kind'] in ('arrival_note','packing_detail'))), next(iter(sheets)))
        if selected not in sheets: raise ValueError('工作表不存在，请重新选择')
        rows, merged = sheets[selected]['rows'], sheets[selected]['merged']
        info = header_info(rows, app, int(source.get('header') or 0))
        mapping = dict(info['mapping'])
        for key, col in (source.get('mapping') or {}).items():
            if key not in LABELS or not isinstance(col, int) or not -1<=col<200: raise ValueError('字段映射无效')
            mapping[key] = col
        kind, reason = classify(selected, rows, info, app)
        if role == 'auto': role = 'arrival_note' if kind == 'arrival_note' else 'packing_detail'
        selection = dict(input=index, sheet=selected, role=role, kind=kind, header_row=info['header_row'], headers=info['headers'], mapping=mapping,
                         sample=[[app.clean(c) if not isinstance(c,dict) else '[公式]' for c in row] for row in rows[:20]])
        selections.append(selection)
        prefix = f'来源{index+1} / {selected}'
        if any('DEMO-ONLY' in app.clean(cell) for row in rows for cell in row) and not app.demo_mode():
            errors.append('演示数据只能在演示模式中使用')
        if kind in ('shipment', 'document_only'):
            errors.append(f'{prefix}：{reason}，不能直接生成库存；请选择物资明细表')
            continue
        if kind in ('unknown', 'material_summary') and not source.get('scope_confirmed'):
            errors.append(f'{prefix}：请确认选定的是入库物资明细，不与汇总重复')
        if mapping['name'] < 0 and mapping['qty'] < 0:
            errors.append(f'{prefix}：未识别名称和数量，请指定表头行及对应列')
            continue
        if sum(c['input']==index and c['kind'] in ('packing_detail','material_summary') for c in candidates)>1:
            warnings.append(f'{prefix}：文件含多个明细/汇总页；本次仅处理选中的工作表，请勿重复入库')
        # Page headings can contain a single explicit box. Never silently carry a blank box.
        page_box = ''
        page_box_evidence = ''
        for row in rows[:info['end']]:
            text = ' '.join(app.clean(v) for v in row if not isinstance(v,dict))
            matches = re.findall(BOX_PATTERN, text)
            if len(matches)==1: page_box, page_box_evidence = matches[0], text
        current_mapping = mapping
        current_headers = info['headers']
        signature_area = False
        for row_index in range(info['end'], len(rows)):
            cells = rows[row_index]
            if not any(app.clean(c) for c in cells): continue
            heading = ' '.join(app.clean(c) for c in cells if not isinstance(c,dict))
            compact = [re.sub(r'\s+', '', app.clean(c)) for c in cells if not isinstance(c,dict)]
            layout_reason = ''
            repeated = header_info([cells], app)
            if repeated['mapping']['name'] >= 0 and repeated['mapping']['qty'] >= 0:
                if not source.get('mapping'): current_mapping = repeated['mapping']
                current_headers = repeated['headers']
                signature_area = False
                layout_reason = '重复表头'
            elif any(c in ('装箱清单','装箱单','箱单') for c in compact):
                page_box = page_box_evidence = ''
                signature_area = False
                layout_reason = '装箱页标题'
            elif sum(bool(re.match(r'^(项目号|项目名称|合同编号|箱件描述|装箱单位|包装类型|长\*宽\*高|净重/毛重)[:：]', c)) for c in compact) >= 2:
                layout_reason = '装箱页基本信息'
            elif sum(c in ('驻场代表','检查','成套','包装','包装日期','验收人','签字','签名') for c in compact) >= 3:
                signature_area = True
                layout_reason = '签字栏标题'
            elif any(c.startswith('本单打印时间') for c in compact):
                layout_reason = '打印信息'
            box_matches = re.findall(BOX_PATTERN, heading)
            if box_matches:
                page_box = box_matches[0] if len(box_matches)==1 else ''
                page_box_evidence = heading
                layout_reason = '完整发运箱号 / 页码'
            identity = f'{index}:{selected}:{row_index+1}'
            patch = corrections.get(identity, {})
            if not isinstance(patch,dict) or set(patch)-EDITABLE: raise ValueError('不支持的人工修正字段')
            material_patch = any(key != 'exclude_reason' for key in patch)
            if layout_reason and not material_patch:
                layout_rows.append(dict(id=identity, input=index, sheet=selected, source_row=row_index+1, reason=layout_reason, raw=heading))
                continue
            record = {key: '' for key in LABELS}
            formula_fields = []
            for key, col in current_mapping.items():
                value = cells[col] if 0<=col<len(cells) else ''
                if isinstance(value,dict): formula_fields.append(key)
                else: record[key] = app.clean(value)
            if current_headers:
                drawing_columns=[i for i,h in enumerate(current_headers) if app.norm(h) in {app.norm(x) for x in ('图号','零件图号','图纸编号','drawing no','drawing number')}]
                drawing_column=next((i for i in drawing_columns if i<len(cells) and cells[i] and not isinstance(cells[i],dict)),-1)
                if drawing_column>=0:
                    record['drawing']=app.clean(cells[drawing_column])
                    if current_mapping.get('spec',-1) in drawing_columns:
                        record['spec']=''
            attr_columns = [i for i,h in enumerate(current_headers) if app.norm(h) in {app.norm(a) for a in app.ALIASES['attribute']}]
            explicit_attribute=current_mapping.get('attribute',-1)
            if explicit_attribute>=0 and explicit_attribute not in attr_columns: attr_columns.append(explicit_attribute)
            if attr_columns:
                record['attribute'] = '；'.join(dict.fromkeys(app.clean(cells[i]) for i in attr_columns if i<len(cells) and cells[i] and not isinstance(cells[i],dict)))
            if signature_area and not material_patch and not record['code'] and not record['unit'] and re.match(r'^\d{4}[/.-]\d{1,2}[/.-]\d{1,2}(?:\s|$)', record['qty']):
                layout_rows.append(dict(id=identity, input=index, sheet=selected, source_row=row_index+1, reason='签字姓名 / 包装日期', raw=heading))
                continue
            if not any(record[k] for k in ('name','qty','code')) and not formula_fields: continue
            if re.match(r'^(合计|总计|小计|签字|签名|编制|审核|检验员|发货人|收货人|第\s*\d+\s*页)', heading.strip()):
                continue
            original = dict(record)
            suggestions, evidence = {}, {}
            if not record['box']:
                col = current_mapping['box']
                for top,bottom,left,right in merged:
                    if top<=row_index+1<=bottom and left<=col+1<=right:
                        v = rows[top-1][left-1]
                        if v and not isinstance(v,dict): suggestions['box']=app.clean(v); evidence['box']=f'合并单元格，第{top}行'
                if not suggestions and page_box: suggestions['box']=page_box; evidence['box']=page_box_evidence
            if page_box and record['box'] != page_box:
                suggestions['box']=page_box
                evidence['box']=page_box_evidence
            # Explicit page metadata is usable evidence; unrelated conflicting box IDs still need review.
            auto_box = bool(page_box and record['box'] != page_box and
                            (not record['box'] or page_box.startswith(record['box']+'-') or page_box in shipment_boxes))
            if auto_box:
                record['box']=page_box
                suggestions.pop('box',None)
            if record['qty'] and not record['unit']:
                match = re.fullmatch(r'([0-9]+(?:\.[0-9]+)?)\s*([^\d\s.]+)', record['qty'])
                if match: suggestions.update(qty=match[1], unit=match[2]); evidence['unit']='数量单元格中的单位，需确认'
            for key, value in patch.items():
                if not isinstance(value,(str,int,float)): raise ValueError('人工修正值必须是文字或数字')
                if len(str(value))>2000: raise ValueError('人工修正内容过长')
                record[key]=app.clean(value)
            if not record['spec'] and not patch.get('spec'):
                record['spec']=inferred_spec(record['name'])
            for key in ('box','unit','doc_no'):
                if not record[key]: record[key]=app.clean(source.get('defaults',{}).get(key) or (p.get('defaults') or {}).get(key))
            record.update(id=identity, input=index, role=role, source=source['source'], sheet=selected, source_row=row_index+1,
                          original=original, raw=[app.clean(c) if not isinstance(c,dict) else '[公式]' for c in cells],
                          suggestions=suggestions, evidence=evidence, page_box=page_box, auto_box=auto_box and 'box' not in patch, corrected=list(patch), issues=[])
            if record.get('exclude_reason'):
                record['status']='已人工排除'; source_rows.append(record); continue
            for field in formula_fields:
                if field not in patch: record['issues'].append(f'{LABELS[field]}为公式，请转为值或人工填写')
            for key in ('name','qty','unit'):
                if not record[key]: record['issues'].append(f'缺少{LABELS[key]}')
            try: record['amount']=app.units(record['qty'])
            except (ValueError,TypeError): record['issues'].append('数量须为大于0的数值')
            if role=='packing_detail' and not record['box']: record['issues'].append('缺少箱号，建议值须人工接受')
            elif role=='packing_detail' and page_box and record['box']!=page_box and 'box' not in patch:
                record['issues'].append(f'表内箱号 {record["box"]} 与页眉发运箱号 {page_box} 不同，请确认完整箱号')
            record['status']='待确认' if record['issues'] else ('人工确认' if patch or source['source']=='manual' else '已识别')
            source_rows.append(record)
    if len(source_rows)>MAX_ROWS: raise ValueError('最多10000条明细')
    known = {r['id'] for r in source_rows+layout_rows}
    if set(corrections)-known: errors.append('来源内容已变动，部分修正不再对应原行；请清空修正后重新识别')
    active = [r for r in source_rows if not r.get('exclude_reason')]
    arrivals = [r for r in active if r['role']=='arrival_note']
    packing = [r for r in active if r['role']=='packing_detail']
    reconciliation = []
    if arrivals and packing:
        totals = {}
        for r in active:
            if not r['doc_no']: r['issues'].append('缺到货单号，无法关联两张单据')
            key = (r['doc_no'],app.norm(r['name']),app.norm(r['spec']),app.norm(r['unit']))
            pair = totals.setdefault(key, dict(doc_no=r['doc_no'], name=r['name'], spec=r['spec'], unit=r['unit'], arrival=0, packing=0, arrival_present=False, packing_present=False))
            side='arrival' if r['role']=='arrival_note' else 'packing'
            pair[side]+=r.get('amount',0); pair[side+'_present']=True
        for pair in totals.values():
            pair['ok']=bool(pair['doc_no']) and pair['arrival_present'] and pair['packing_present'] and pair['arrival']==pair['packing']
            if not pair['ok']: errors.append(f"{pair['doc_no'] or '单号缺失'} / {pair['name']}：到货单与装箱单数量不一致或缺对应明细，请核对来源后修正")
            pair['arrival']=app.number(pair['arrival']); pair['packing']=app.number(pair['packing'])
            reconciliation.append(pair)
    elif arrivals:
        packing=arrivals
        for r in arrivals:
            if not inputs[r['input']].get('confirm_detail'): r['issues'].append('仅有到货单，请明确确认按这些明细入库')
            if not r['box']: r['issues'].append('按到货单入库仍需补充箱号')
    for r in active:
        if r['issues']: r['status']='待确认'
        errors.extend(f"来源{r['input']+1} / {r['sheet']} 第{r['source_row']}行：{issue}" for issue in r['issues'])
    valid = [r for r in packing if not r['issues']]
    # Reuse the existing ledger parser for identity, item codes, positions and aggregation.
    csv_text=io.StringIO(); writer=csv.writer(csv_text)
    writer.writerow([app.ALIASES[k][0] for k in app.ALIASES])
    defaults=dict(p.get('defaults') or {})
    for r in valid:
        canonical=dict(r)
        if r.get('drawing'):
            canonical['attribute']='；'.join(filter(None,(r.get('attribute',''),f"图号：{r['drawing']}")))
        writer.writerow([canonical.get(k,defaults.get(k,'')) for k in app.ALIASES])
    canonical=dict(p, filename='intake.csv', content=base64.b64encode(csv_text.getvalue().encode('utf-8')).decode(), header=1, sheet='CSV')
    canonical.pop('mapping',None)
    result,_=app.parse_file(canonical)
    errors.extend(result['errors']); warnings.extend(result['warnings'])
    if not active: errors.append('没有可入库明细，请选择工作表、调整字段映射或使用手动录入')
    for r in result['rows']:
        original_index=int(r['source_row'])-2
        if 0<=original_index<len(valid): r['source_row']=valid[original_index]['source_row']
    selection_identity=[dict(input=s['input'],sheet=s['sheet'],role=s['role']) for s in selections]
    # Automatic and explicit choices share the legacy resolved-role identity.
    identity_raw=packed({'format':'warehouse-intake-v1','sources':[{**source,'role':selection['role']} for source,selection in zip(originals,selections)]})
    result['hash']=hashlib.sha256(identity_raw+b'\0'+packed(selection_identity)).hexdigest()
    selection_hash=hashlib.sha256(raw+packed({'inputs':inputs,'corrections':corrections,'defaults':defaults})).hexdigest()
    with app.connect() as c:
        if c.execute('SELECT 1 FROM documents WHERE file_hash=?',(result['hash'],)).fetchone(): errors.append('此来源及工作表已导入，请勿重复入库')
    result.update(source_rows=source_rows, layout_rows=layout_rows, candidates=candidates, selections=selections, reconciliation=reconciliation,
                  sheet=' / '.join(s['sheet'] for s in selections), sheets=[s['sheet'] for s in selections], sheet_kind='packing',
                  selection_hash=selection_hash, errors=list(dict.fromkeys(errors))[:100], warnings=list(dict.fromkeys(warnings))[:100])
    return result, raw, '入库来源.intake.json'
