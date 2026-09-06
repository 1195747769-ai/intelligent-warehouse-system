import base64, concurrent.futures, importlib, io, os, tempfile, unittest, uuid
import openpyxl
import server as app

class LedgerTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); app.DATA=app.Path(self.tmp.name); app.init()
    def tearDown(self): self.tmp.cleanup()
    def request(self,**kw): return dict(operator='TEST',party='TEST CREW',purpose='Synthetic test',request_key=str(uuid.uuid4()),**kw)
    def receipt(self,ref='R1',qty='98',code='M1',state='AVAILABLE',kind='IN'):
        content=f'物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态\n{code},Test material,M12,EA,{qty},B1,C01,WH,A1,{state}\n'
        return self.request(external_ref=ref,filename=ref+'.csv',content=base64.b64encode(content.encode()).decode(),kind=kind)
    def pos(self,state='AVAILABLE'):
        return next(s for s in app.overview()['stock'] if s['state']==state)
    def test_complete_flow(self):
        app.import_post(self.receipt('design','1000',kind='BASELINE'))
        app.change_design(dict(self.request(),code='M1',qty='100',reason='Approved test change'))
        incoming=app.import_post(self.receipt())
        p=self.request(kind='STATE',state='PENDING',lines=[{'stock_id':self.pos()['id'],'qty':'3'}]); app.stock_post(p)
        self.assertEqual(self.pos()['qty'],95)
        with self.assertRaises(ValueError): app.stock_post(self.request(kind='OUT',lines=[{'stock_id':self.pos('PENDING')['id'],'qty':1}]))
        app.stock_post(self.request(kind='STATE',state='SPARE',lines=[{'stock_id':self.pos()['id'],'qty':10}]))
        issue=app.stock_post(self.request(kind='OUT',lines=[{'stock_id':self.pos()['id'],'qty':20}]))
        app.stock_post(self.request(kind='HANDOVER',lines=[{'stock_id':self.pos('SPARE')['id'],'qty':10}]))
        data=app.overview(); item=data['items'][0]
        self.assertEqual((item['current'],item['received'],item['issued'],item['handed'],item['onhand'],item['outstanding']),(1100,98,20,10,68,1002))
        self.assertEqual(self.pos()['qty'],65)
        app.doc_detail(issue['id']); app.doc_detail(issue['id'])
        self.assertEqual(self.pos()['qty'],65)
        app.reverse(dict(self.request(),id=issue['id']))
        self.assertEqual(self.pos()['qty'],85)
        with self.assertRaises(ValueError):app.reverse(dict(self.request(),id=incoming['id']))
    def test_transaction_ledger_matches_movements(self):
        app.import_post(self.receipt('LEDGER',qty='10'))
        data=app.overview()
        self.assertEqual(len(data['ledger']),1)
        self.assertTrue(all(row['kind']=='IN' and row['delta']>0 for row in data['ledger']))
        self.assertEqual({row['number'] for row in data['ledger']},{data['documents'][0]['number']})

    def test_server_balance_fields_and_reconciliation(self):
        app.import_post(self.receipt('BALANCE',qty='8'))
        before=app.overview(); stock=before['stock'][0]
        self.assertEqual((stock['received_total'],stock['available_qty']),(8,8))
        app.stock_post(self.request(kind='OUT',lines=[{'stock_id':stock['id'],'qty':1}]))
        after=app.overview(); stock=after['stock'][0]
        self.assertEqual((stock['received_total'],stock['available_qty']),(8,7))
        self.assertTrue(after['reconciliation']['ok'])

    def test_password_protected_admin_delete_rolls_back_and_logs(self):
        incoming=app.import_post(self.receipt('ADMIN-DELETE',qty='8'))
        with self.assertRaisesRegex(ValueError,'管理员密码'):
            app.admin_delete(dict(id=incoming['id'],password='wrong',purpose='mistake',operator='ADMIN'))
        app.admin_password_set({'new_password':'secret1'})
        with self.assertRaisesRegex(ValueError,'管理员密码错误'):
            app.admin_delete(dict(id=incoming['id'],password='wrong',purpose='mistake',operator='ADMIN'))
        app.admin_delete(dict(id=incoming['id'],password='secret1',purpose='误导入测试',operator='ADMIN'))
        self.assertFalse(app.overview()['documents'])
        self.assertFalse(app.overview()['items'])
        with app.connect() as c:
            action=c.execute('SELECT action,document_number,reason,operator FROM admin_actions').fetchone()
            self.assertEqual(tuple(action),('DELETE', 'IN-'+action['document_number'].split('-',1)[1], '误导入测试', 'ADMIN'))

    def test_admin_delete_rejects_receipt_already_used_by_issue(self):
        incoming=app.import_post(self.receipt('ADMIN-BLOCK',qty='8'))
        stock=self.pos()
        app.stock_post(self.request(kind='OUT',lines=[{'stock_id':stock['id'],'qty':'1'}]))
        app.admin_password_set({'new_password':'secret1'})
        with self.assertRaisesRegex(ValueError,'后续库存'):
            app.admin_delete(dict(id=incoming['id'],password='secret1',purpose='误导入',operator='ADMIN'))

    def test_receipt_cannot_exceed_known_design_quantity(self):
        app.import_post(self.receipt('BASELINE',qty='5',kind='BASELINE'))
        with self.assertRaisesRegex(ValueError,'超过合同/设计数量'):
            app.import_post(self.receipt('OVER',qty='6'))
        self.assertEqual(len(app.overview()['documents']),1)
    def test_duplicates_and_transactions(self):
        p=self.receipt(); d=app.import_post(p)
        self.assertEqual(app.import_post(p),d)
        p['request_key']=str(uuid.uuid4())
        with self.assertRaises(ValueError):app.import_post(p)
        p=self.receipt(qty='99')
        with self.assertRaises(app.sqlite3.IntegrityError):app.import_post(p)
        before=len(app.overview()['documents'])
        with self.assertRaises(ValueError):app.stock_post(self.request(lines=[{'stock_id':self.pos()['id'],'qty':10},{'stock_id':self.pos()['id'],'qty':1000}]))
        self.assertEqual(self.pos()['qty'],98)
        self.assertEqual(len(app.overview()['documents']),before)
    def test_repeated_rows_aggregate(self):
        content='物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态\nM1,Test material,M12,EA,2,B1,C01,WH,A1,AVAILABLE\nM1,Test material,M12,EA,3,B1,C01,WH,A1,AVAILABLE\n'
        p=self.request(filename='repeat.csv',content=base64.b64encode(content.encode()).decode(),external_ref='REPEAT')
        result,_=app.parse_file(p)
        self.assertEqual(result['errors'],[]);self.assertEqual(len(result['rows']),1);self.assertEqual(float(result['rows'][0]['qty']),5)
        app.import_post(p);self.assertEqual(self.pos()['qty'],5)
    def test_contract_inventory_group(self):
        content='物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态,合同\nM1,组件支架,M12,EA,8,PV25090501,C01,WH,A1,AVAILABLE,7300002797-组件\n'
        p=self.request(filename='contract.csv',content=base64.b64encode(content.encode()).decode(),external_ref='CONTRACT-1',arrival_date='2025-09-05')
        receipt=app.import_post(p); self.assertEqual(app.doc_detail(receipt['id'])['arrival_date'],'2025-09-05'); sid=self.pos()['id']
        app.stock_post(self.request(kind='OUT',lines=[{'stock_id':sid,'qty':3}]))
        group=next(g for g in app.overview()['contracts'] if g['code']=='7300002797-组件')
        self.assertEqual((group['received'],group['issued'],group['onhand']), (8,3,5))
        item=group['items'][0]
        self.assertTrue(item['inbound_first']); self.assertTrue(item['outbound_last'])
    def test_boxes_are_distinct_positions(self):
        content='物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态\nM1,Test material,M12,EA,2,B1,C01,WH,A1,AVAILABLE\nM1,Test material,M12,EA,3,B1,C02,WH,A1,AVAILABLE\n'
        p=self.request(filename='boxes.csv',content=base64.b64encode(content.encode()).decode(),external_ref='BOXES')
        app.import_post(p); positions=[s for s in app.overview()['stock'] if s['code']=='M1']; self.assertEqual(sorted((s['box'],s['qty']) for s in positions),[('C01',2),('C02',3)])
    def test_merged_box_value_is_carried_down(self):
        content='物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态\nM1,Test material,M12,EA,2,B1,C01,WH,A1,AVAILABLE\nM2,Second material,M13,EA,3,B1,,WH,A1,AVAILABLE\n'
        p=self.request(filename='merged-box.csv',content=base64.b64encode(content.encode()).decode(),external_ref='MERGED-BOX')
        result,_=app.parse_file(p); self.assertEqual(result['errors'],[]); self.assertTrue(all(r['box']=='C01' for r in result['rows']))
    def test_multiple_attribute_columns_are_kept(self):
        content='物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态,装配描述,备注\nM1,Test material,M12,EA,2,B1,C01,WH,A1,AVAILABLE,成套设备,厂家序列号SN-01\n'
        p=self.request(filename='attributes.csv',content=base64.b64encode(content.encode()).decode(),external_ref='ATTRIBUTES')
        result,_=app.parse_file(p); self.assertEqual(result['errors'],[]); self.assertEqual(result['rows'][0]['attribute'],'成套设备；厂家序列号SN-01')
        app.import_post(p); self.assertEqual(app.overview()['stock'][0]['attribute'],'成套设备；厂家序列号SN-01')
    def test_box_remark_becomes_child_contents_without_extra_stock(self):
        note='（拼接钢板-轨道下部 AX300A-0300-0004 64件\n计数拨片1 AX300A-0300-0001 8个）'
        content='物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态,备注\nM1,主变室辅料,,箱,1,B1,C01,WH,A1,AVAILABLE,"'+note+'"\n'
        p=self.request(filename='box-contents.csv',content=base64.b64encode(content.encode()).decode(),external_ref='BOX-CONTENTS')
        result,_=app.parse_file(p); children=result['rows'][0]['contents']; self.assertEqual([(x['code'],x['name'],x['qty'],x['unit']) for x in children],[('AX300A-0300-0004','拼接钢板-轨道下部','64','件'),('AX300A-0300-0001','计数拨片1','8','个')])
        app.import_post(p); stock=app.overview()['stock'][0]; self.assertEqual(stock['qty'],1); self.assertEqual(len(stock['contents']),2)
    def test_shipment_sheet_cannot_post_receipt(self):
        content='箱件号,货物名称,包装类型,单位,数量\nC01,Test material,木箱,箱,1\n'
        p=self.request(filename='shipment.csv',content=base64.b64encode(content.encode()).decode(),external_ref='SHIPMENT',defaults={'batch':'B1','warehouse':'WH','bin':'A1'})
        with self.assertRaisesRegex(ValueError,'发货清单'):
            app.parse_file(p)
        with self.assertRaisesRegex(ValueError,'发货清单'):
            app.import_post(p)

    def test_generic_shipment_headers_are_classified_before_receipt(self):
        content='箱号,货物名称,包装类型,单位,数量\nC01,Test material,木箱,箱,1\n'
        p=self.request(filename='generic.csv',content=base64.b64encode(content.encode()).decode(),external_ref='GENERIC-SHIPMENT')
        with self.assertRaisesRegex(ValueError,'发货清单'):
            app.parse_file(p)
    def test_packing_sheet_name_wins_over_workbook_filename(self):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='一期装箱清单';ws.append(['物资编码','物资名称','规格型号','单位','数量','箱号']);ws.append(['M1','Test material','M12','EA',2,'C01'])
        stream=io.BytesIO();wb.save(stream)
        p=self.request(filename='阜康发货清单.xlsx',content=base64.b64encode(stream.getvalue()).decode(),external_ref='PACKING-SHEET',defaults={'batch':'B1','warehouse':'WH','bin':'A1'},sheet='一期装箱清单')
        result=app.import_post(p);self.assertTrue(result['id'])
    def test_contract_number_is_read_from_workbook_cover(self):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='一期装箱清单'
        ws['H3']='合同编号：7300002797'
        ws.append([]);ws.append([]);ws.append([]);ws.append([])
        ws.append(['','箱号','序号','物料编码','','零件图号','','零件名称','','数量','','单位','装配描述','备注'])
        ws.append(['','C01','1','','','','','Test material','','2','','件','',''])
        stream=io.BytesIO();wb.save(stream)
        p=self.request(filename='cover-contract.xlsx',content=base64.b64encode(stream.getvalue()).decode(),external_ref='COVER-CONTRACT',defaults={'batch':'B1','warehouse':'WH','bin':'A1'},sheet='一期装箱清单')
        result,_=app.parse_file(p)
        self.assertEqual(result['errors'],[])
        self.assertEqual({r['package'] for r in result['rows']},{'7300002797'})
        self.assertTrue(any('7300002797' in w for w in result['warnings']))
    def test_fixed_segments_and_selected_segment_override_cover_contract(self):
        self.assertEqual([code for code, _, _ in app.SEGMENTS], [f'A{i}' for i in range(1,17)])
        content='物资编码,物资名称,规格型号,单位,数量,箱号,合同\nM1,Test material,M12,EA,2,C01,7300002797\n'
        p=self.request(filename='segment.csv',content=base64.b64encode(content.encode()).decode(),kind='IN',segment_required=True,arrival_date='2026-09-06',defaults={'package':'A1','batch':'B01','warehouse':'WH','bin':'A1','state':'AVAILABLE'})
        result,_=app.parse_file(p)
        self.assertEqual(result['errors'],[])
        self.assertEqual({r['package'] for r in result['rows']},{'A1'})
        self.assertEqual({r['batch'] for r in result['rows']},{'B01'})
    def test_receipt_defaults_are_applied_once(self):
        content='物资编码,物资名称,规格型号,单位,数量,箱号\n,Test material,M12,EA,2,C01\n,Test material,M12,EA,3,C02\n'
        p=self.request(filename='minimal.csv',content=base64.b64encode(content.encode()).decode(),external_ref='MINIMAL',arrival_date='2025-09-05',batch_device='组件',batch_number='1')
        result,_=app.parse_file(p)
        self.assertEqual(result['errors'],[])
        self.assertEqual((result['rows'][0]['batch'],result['rows'][0]['code'],result['rows'][0]['warehouse'],result['rows'][0]['bin']),( '组件-250905-01','组件-250905-01-001','主仓库','待分配'))
        self.assertEqual({r['code'] for r in result['rows']},{'组件-250905-01-001'})

    def test_code_less_item_reuses_internal_code_across_batches(self):
        def payload(ref,batch,qty):
            content=f'物资编码,物资名称,规格型号,单位,数量,箱号\n,Same item,M12,EA,{qty},C01\n'
            return self.request(filename=ref+'.csv',content=base64.b64encode(content.encode()).decode(),external_ref=ref,segment_required=True,defaults={'package':'A1','batch':batch,'warehouse':'WH','bin':'A1','state':'AVAILABLE'})
        first,_=app.parse_file(payload('NO-CODE-1','B01',2)); app.import_post(payload('NO-CODE-1','B01',2))
        second,_=app.parse_file(payload('NO-CODE-2','B02',3))
        self.assertEqual(first['rows'][0]['code'],second['rows'][0]['code'])
        self.assertTrue(any('沿用了既有内部编码' in w for w in second['warnings']))
    def test_business_reference_can_be_blank(self):
        p=self.receipt('TEMP')
        p['external_ref']=''
        result=app.import_post(p)
        doc=app.doc_detail(result['id'])
        self.assertEqual(doc['external_ref'],doc['number'])
    def test_concurrent_issue(self):
        app.import_post(self.receipt(qty='10')); sid=self.pos()['id']
        def issue(_):
            try:app.stock_post(self.request(lines=[{'stock_id':sid,'qty':7}]));return True
            except ValueError:return False
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as ex:result=list(ex.map(issue,range(2)))
        self.assertEqual(sorted(result),[False,True]);self.assertEqual(self.pos()['qty'],3)
    def test_units_moves_and_baseline_lock(self):
        app.import_post(self.receipt(qty='.123456'))
        app.stock_post(self.request(kind='MOVE',warehouse='WH2',bin='B2',lines=[{'stock_id':self.pos()['id'],'qty':'.1'}]))
        self.assertAlmostEqual(sum(s['qty'] for s in app.overview()['stock']),.123456, places=6)
        with self.assertRaises(ValueError):app.units('.0000001')
        p=self.receipt('new',qty=1); p['content']=base64.b64encode(base64.b64decode(p['content']).replace(b',EA,',b',KG,')).decode()
        with self.assertRaises(ValueError):app.import_post(p)
        app.import_post(self.receipt('base',qty=10,kind='BASELINE'))
        with self.assertRaises(ValueError):app.import_post(self.receipt('base2',qty=20,kind='BASELINE'))
    def test_excel_mapping_and_formulas(self):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='Packing'
        ws.append(['Supplier title']);ws.append(['SKU','Description','Specification','UOM','Qty']);ws.append(['X1','Part','M12','EA',20])
        stream=io.BytesIO();wb.save(stream)
        p=self.request(filename='supplier.xlsx',content=base64.b64encode(stream.getvalue()).decode(),header=2,defaults={'batch':'B','box':'C01','warehouse':'W','bin':'A'},external_ref='XLSX')
        result,_=app.parse_file(p);self.assertEqual(result['errors'],[]);self.assertEqual(result['rows'][0]['qty'],'20');app.import_post(p)
        ws.cell(3,5,'=10+10');stream=io.BytesIO();wb.save(stream);p['content']=base64.b64encode(stream.getvalue()).decode()
        result,_=app.parse_file(p);self.assertTrue(any('公式' in e for e in result['errors']))

if __name__=='__main__':unittest.main(verbosity=2)
