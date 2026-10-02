import base64
import io
import os
import tempfile
import unittest
import uuid
import json
import zipfile
from unittest.mock import patch
from pathlib import Path

import openpyxl
import server as app


class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_data = app.DATA
        app.DATA = Path(self.tmp.name)
        app.init()

    def tearDown(self):
        app.DATA = self.old_data
        self.tmp.cleanup()

    def request(self, inputs, **extra):
        return dict(inputs=inputs, kind='IN', operator='TEST', party='虚构厂家',
                    defaults=dict(package='A1', batch='TEST-B01', warehouse='测试仓', bin='01', state='AVAILABLE'),
                    segment_required=True, request_key=str(uuid.uuid4()), **extra)

    def paste(self, text, role='packing_detail'):
        return dict(source='paste', role=role, text=text)

    def test_auto_source_and_manual_arrival_confirmation(self):
        source=self.paste('箱号\t物资名称\t数量\t单位\n7-4\t轴承\t2\t个','auto')
        result,_,_=app.parse_intake(self.request([source]))
        self.assertEqual(result['selections'][0]['role'],'packing_detail')
        self.assertEqual(result['errors'],[])
        source=self.paste('到货单\n箱号\t物资名称\t数量\t单位\n7-4\t轴承\t2\t个','auto')
        result,_,_=app.parse_intake(self.request([source]))
        self.assertEqual(result['selections'][0]['role'],'arrival_note')
        self.assertTrue(result['errors'])
        source['confirm_detail']=True
        result,_,_=app.parse_intake(self.request([source]))
        self.assertEqual(result['errors'],[])
        self.assertEqual(app.overview()['stock'],[])

    def test_auto_and_explicit_roles_share_existing_duplicate_identity(self):
        text='箱号\t物资名称\t数量\t单位\n7-4\t轴承\t2\t个'
        explicit=self.request([self.paste(text)])
        old_result,old_raw,_=app.parse_intake(explicit)
        import hashlib
        import intake
        identity=[dict(input=s['input'],sheet=s['sheet'],role=s['role']) for s in old_result['selections']]
        legacy_hash=hashlib.sha256(old_raw+b'\0'+intake.packed(identity)).hexdigest()
        self.assertEqual(old_result['hash'],legacy_hash)
        app.import_post(explicit)
        auto=self.request([self.paste(text,'auto')])
        result,_,_=app.parse_intake(auto)
        self.assertEqual(result['hash'],legacy_hash)
        self.assertTrue(any('已导入' in error for error in result['errors']))
        with self.assertRaises(ValueError): app.import_post(auto)
        self.assertEqual(len(app.overview()['documents']),1)

    def test_local_ai_only_returns_grounded_suggestions_without_posting(self):
        import local_ai
        p=self.request([self.paste('箱号\t物资名称\t数量\t单位\n7-4\t圆螺母M90x6 锻钢34Cr2Ni2Mo\t40\t个')])
        result,_,_=app.parse_intake(p); row=result['source_rows'][0]
        fake={'rows':[{'id':row['id'],'fields':{'name':'圆螺母','spec':'不存在的型号','qty':'999','unit':'箱','box':'另一个箱'}}], 'mappings':[]}
        with patch.object(local_ai,'infer',return_value=fake): suggestion=app.intake_ai_post(p)
        self.assertEqual(suggestion['suggestions'],[{'id':row['id'],'fields':{'name':'圆螺母'}}])
        self.assertTrue(suggestion['requires_confirmation'])
        self.assertEqual(app.overview()['stock'],[])
        self.assertEqual(app.overview()['documents'],[])
        self.assertEqual(app.recognition_list({})['total'],0)
        p['corrections']={row['id']:{'name':'已人工确认的名称'}}
        with patch.object(local_ai,'infer',return_value=fake): self.assertEqual(app.intake_ai_post(p)['suggestions'],[])

    def test_local_ai_rejects_box_number_invented_as_spec_and_duplicate_columns(self):
        import local_ai
        row=dict(name='钢板06Cr19Ni10',spec='',drawing='',raw=['7-4','钢板06Cr19Ni10','2','件'])
        self.assertEqual(local_ai.grounded_fields(row,{'spec':'7-4 06Cr19Ni10'}),{})
        result={'source_rows':[],'selections':[dict(input=0,mapping={'name':-1,'qty':-1,'unit':-1},sample=[['名称','数量','单位']])]}
        with patch.object(local_ai,'infer',return_value={'rows':[],'mappings':[dict(input=0,header=1,mapping={'name':0,'qty':0,'unit':0})]}):
            self.assertEqual(local_ai.suggest(result)['mappings'],[])
        with patch.object(local_ai,'infer',return_value={'rows':[],'mappings':[dict(input=0,header=1,mapping={'name':0,'qty':1,'unit':2})]}):
            self.assertEqual(local_ai.suggest(result)['mappings'][0]['mapping'],{'name':0,'qty':1,'unit':2})

    def test_real_sample_model_suffixes_page_box_and_ai_drawing(self):
        import intake
        import local_ai
        for original, expected in [('球阀 Q47F-100P 凸面 DN50','Q47F-100P DN50'),
                                   ('密封圈 KDS10X','KDS10X'),
                                   ('焊条 AWS A5.4 E309MoL-16 CHS042 Φ3.2','AWS A5.4 E309MoL-16 CHS042 Φ3.2')]:
            self.assertEqual(intake.inferred_spec(original),expected)
        source=self.paste('装箱清单\n箱  件  号：\t\t7-2\n箱号\t物资名称\t数量\t单位\t零件图号\n7\t主轴加工图\t1\t个\t1S19796')
        preview,_,_=app.parse_intake(self.request([source]))
        row=preview['source_rows'][0]
        self.assertEqual((row['box'],row['drawing'],preview['errors']),('7-2','1S19796',[]))
        self.assertEqual(local_ai.grounded_fields(row,{'drawing':'主轴加工图'}),{})
        row['drawing']=''
        self.assertEqual(local_ai.grounded_fields(row,{'drawing':'主轴加工图'}),{})
        row['raw'].append('图号：1S19796')
        self.assertEqual(local_ai.grounded_fields(row,{'drawing':'1S19796'}),{'drawing':'1S19796'})

    def test_upgraded_ai_checks_complex_rows_and_more_than_six(self):
        import local_ai
        rows=[dict(id=str(i),name='钢板06Cr19Ni10' if i>=20 else '主轴',spec='06Cr19Ni10' if i>=20 else '',issues=[],raw=[],exclude_reason='') for i in range(40)]
        rows[-1]['issues']=['待核对']
        rows[38]['exclude_reason']='人工移除'
        def answer(data):
            self.assertEqual(len(data['rows']),32)
            self.assertEqual([r['id'] for r in data['rows'][:2]],['39','20'])
            self.assertNotIn('38',[r['id'] for r in data['rows']])
            return {'rows':[dict(id=str(i),fields={'name':'钢板'}) for i in range(20,28)],'mappings':[]}
        with patch.object(local_ai,'infer',side_effect=answer):
            result=local_ai.suggest({'source_rows':rows,'selections':[]})
        self.assertEqual((result['checked'],len(result['suggestions'])),(32,8))
        row=dict(name='圆螺母M90x6 锻钢34Cr2Ni2Mo',spec='M90x6 34Cr2Ni2Mo',drawing='',raw=[])
        self.assertEqual(local_ai.grounded_fields(row,{'spec':'锻钢34Cr2Ni2Mo'}),{})
        self.assertEqual(local_ai.grounded_fields(row,{'spec':'M90x6 锻钢34Cr2Ni2Mo'}),{'spec':'M90x6 锻钢34Cr2Ni2Mo'})

    def test_source_download_returns_original_upload_and_zip(self):
        sources=[dict(source='upload',filename='厂家装箱单.xls',content=base64.b64encode(b'original-xls').decode())]
        raw=json.dumps(dict(format='warehouse-intake-v1',sources=sources)).encode()
        _,rel=app.store_source(raw,'source.intake.json')
        content,mime,name=app.source_download(rel)
        self.assertEqual((content,mime,name),(b'original-xls','application/vnd.ms-excel','厂家装箱单.xls'))
        sources.append(dict(source='upload',filename='厂家到货单.xlsx',content=base64.b64encode(b'original-xlsx').decode()))
        _,rel=app.store_source(json.dumps(dict(format='warehouse-intake-v1',sources=sources)).encode(),'source.intake.json')
        content,mime,name=app.source_download(rel)
        self.assertEqual((mime,name),('application/zip','入库原件.zip'))
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            self.assertEqual(z.read('厂家装箱单.xls'),b'original-xls')
            self.assertEqual(z.read('厂家到货单.xlsx'),b'original-xlsx')

    def test_paste_and_manual_source_download_as_literal_excel(self):
        for source in [self.paste('箱号\t名称\t数量\t单位\n7-4\t=原文\t2\t个'),
                       dict(source='manual',rows=[dict(box='7-4',name='=原文',qty='2',unit='个')])]:
            _,rel=app.store_source(json.dumps(dict(format='warehouse-intake-v1',sources=[source])).encode(),'source.intake.json')
            content,mime,name=app.source_download(rel)
            self.assertTrue(name.endswith('.xlsx'))
            wb=openpyxl.load_workbook(io.BytesIO(content))
            cells=list(wb.active[2]);cell=next(c for c in cells if c.value=='=原文')
            self.assertEqual(cell.data_type,'s')
            wb.close()

    def test_paste_provenance_and_posting(self):
        p = self.request([self.paste('到货单号\t箱号\t物资名称\t数量\t单位\nDEMO-01\tB01\t螺栓\t12\t件')])
        result = app.preview_post(p)
        self.assertEqual(result['candidates'][0]['kind'], 'packing_detail')
        self.assertEqual(result['errors'], [])
        self.assertEqual(result['source_rows'][0]['source_row'], 2)
        self.assertEqual(result['rows'][0]['qty'], '12')
        self.assertEqual(app.overview()['stock'], [])
        p['recognition_id'] = result['recognition_id']
        posted = app.import_post(p)
        self.assertEqual(app.import_post(p), posted)
        self.assertEqual(app.overview()['stock'][0]['qty'], 12)

    def test_reconciliation_blocks_and_correction_is_reparsed(self):
        p = self.request([self.paste('到货单号\t物资名称\t数量\t单位\nDH01\t螺栓\t12\t件', 'arrival_note'),
                          self.paste('到货单号\t箱号\t物资名称\t数量\t单位\nDH01\tB01\t螺栓\t10\t件')])
        r = app.preview_post(p)
        self.assertTrue(any('数量不一致' in e for e in r['errors']))
        with self.assertRaises(ValueError): app.import_post(p)
        row = next(x for x in r['source_rows'] if x['role'] == 'packing_detail')
        p['corrections'] = {row['id']: {'qty': '12'}}
        fixed = app.preview_post(p)
        self.assertEqual(fixed['errors'], [])
        p['recognition_id'] = fixed['recognition_id']
        p['rows'] = [dict(qty='9999')]  # browser preview rows are never trusted
        app.import_post(p)
        self.assertEqual(app.overview()['stock'][0]['qty'], 12)

    def test_manual_missing_box_unit_and_bad_quantity_stay_visible(self):
        p = self.request([dict(source='manual', role='packing_detail', rows=[dict(name='电缆', qty='待核')])])
        r = app.preview_post(p)
        self.assertEqual(len(r['source_rows']), 1)
        self.assertTrue(r['errors'])
        row_id = r['source_rows'][0]['id']
        p['corrections'] = {row_id: dict(qty='5', unit='米', box='B02')}
        self.assertEqual(app.preview_post(p)['errors'], [])

    def test_workbook_sheets_and_explicit_header_box(self):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = '装箱单'
        ws.append(['发运箱号：7-4'])
        ws.append(['名称', '数量', '单位'])
        ws.append(['垫片', 3, '件'])
        wb.create_sheet('发运单').append(['箱号', '毛重', '箱数'])
        mem = io.BytesIO(); wb.save(mem)
        p = self.request([dict(source='upload', role='packing_detail', filename='发货清单.xlsx', content=base64.b64encode(mem.getvalue()).decode())])
        r = app.preview_post(p)
        self.assertEqual(len(r['candidates']), 2)
        row = r['source_rows'][0]
        self.assertEqual(row['box'], '7-4')
        self.assertTrue(row['auto_box'])
        self.assertEqual(row['original']['box'],'')
        self.assertEqual(r['errors'],[])
        p['corrections'] = {row['id']: dict(box='7-4')}
        self.assertEqual(app.preview_post(p)['errors'], [])

    def test_shipment_and_lock_file_blocked(self):
        p = self.request([self.paste('发运单\n箱号\t名称\t箱数\t毛重\nB01\t设备箱\t1\t900')])
        self.assertTrue(app.preview_post(p)['errors'])
        lock = self.request([dict(source='upload', role='packing_detail', filename='~$箱单.xlsx', content='eA==')])
        with self.assertRaisesRegex(ValueError, '锁文件'): app.preview_post(lock)

    def test_confirmed_exclusions_do_not_drop_uncertain_material(self):
        p=self.request([self.paste('箱号\t名称\t数量\t单位\nB01\t零件\t待核\t件\nB02\t本次不要的物资\t待核\t件')])
        r=app.preview_post(p)
        self.assertEqual(len(r['source_rows']),2)
        self.assertTrue(r['errors'])
        p['corrections']={r['source_rows'][0]['id']:dict(qty='2'),r['source_rows'][1]['id']:dict(exclude_reason='签字区')}
        self.assertEqual(app.preview_post(p)['errors'],[])

    def test_paged_packing_filters_layout_but_preserves_uncertain_items(self):
        text=('装箱清单\n发运箱号：7-1\n箱号\t名称\t数量\t单位\t装配描述\t备注\n'
              '7\t护盖\t1\t个\t主轴装配\t核对实物\n'
              '驻场代表\t成套\t包装日期\t\t检查\t验收人\n'
              '\t签字人\t2022/5/19 8:56:43\n'
              '装箱清单\n项目名称：示例\t合同编号：X\n发运箱号：7-4\n'
              '箱号\t名称\t数量\t单位\t装配描述\t备注\n'
              '\t螺栓\t待核\t个\t主轴装配\n')
        p=self.request([self.paste(text)])
        r=app.preview_post(p)
        self.assertEqual(len(r['source_rows']),2)
        self.assertGreaterEqual(len(r['layout_rows']),5)
        first,second=r['source_rows']
        self.assertEqual(first['attribute'],'主轴装配；核对实物')
        self.assertEqual(first['box'],'7-1')
        self.assertEqual(second['box'],'7-4')
        self.assertTrue(any('数量' in x for x in second['issues']))
        p['corrections']={row['id']:dict(box=row['page_box']) for row in r['source_rows']}
        self.assertTrue(app.preview_post(p)['errors'])
        p['corrections'][second['id']]['qty']='3'
        fixed=app.preview_post(p)
        self.assertEqual(fixed['errors'],[])
        self.assertEqual({row['box'] for row in fixed['rows']},{'7-1','7-4'})

    def test_part_drawing_code_stays_separate_from_spec_and_vendor_code(self):
        p=self.request([self.paste('箱号\t物料编码\t零件图号\t零件名称\t数量\t单位\t装配描述\n'
                                   '7-4\t8100.5000312440\t3S31712\t圆螺母M90x6 锻钢34Cr2Ni2Mo\t40\t个\t转子装配')])
        r=app.preview_post(p)
        row=r['source_rows'][0]
        self.assertEqual(row['code'],'8100.5000312440')
        self.assertEqual(row['drawing'],'3S31712')
        self.assertEqual(row['spec'],'M90x6 34Cr2Ni2Mo')
        self.assertNotIn('3S31712',row['spec'])
        self.assertTrue(any('图号：3S31712' in item['attribute'] for item in r['rows']))

    def test_spec_inference_keeps_alloy_grades_and_tonnage(self):
        from intake import inferred_spec
        self.assertEqual(inferred_spec('2x20x2000钢板06Cr19Ni10'),'2x20x2000 06Cr19Ni10')
        self.assertEqual(inferred_spec('圆螺母M90x6 锻钢34Cr2Ni2Mo'),'M90x6 34Cr2Ni2Mo')
        self.assertEqual(inferred_spec('85t卸扣'),'85t')
        self.assertEqual(inferred_spec('主轴护盖'),'')

    def test_distinct_spec_and_drawing_survive_preview_and_posting(self):
        sources=[self.paste('箱号\t物资名称\t规格型号\t图号\t数量\t单位\n'
                            '1\t测试电机\tY100L-4 / IP55\tDWG-01\t1\t台'),
                 dict(source='manual',role='packing_detail',rows=[
                     dict(box='1',name='测试电机',spec='Y100L-4 / IP55',drawing='DWG-01',qty='1',unit='台')])]
        for source in sources:
            with self.subTest(source=source['source']):
                p=self.request([source])
                result=app.preview_post(p)
                self.assertEqual(result['errors'],[])
                self.assertEqual((result['source_rows'][0]['spec'],result['source_rows'][0]['drawing']),
                                 ('Y100L-4 / IP55','DWG-01'))
                self.assertEqual(result['rows'][0]['spec'],'Y100L-4 / IP55')
                receipt=app.import_post(p)
                self.assertEqual(app.doc_detail(receipt['id'])['lines'][0]['spec'],'Y100L-4 / IP55')
                self.assertEqual(app.overview()['stock'][0]['spec'],'Y100L-4 / IP55')

    def test_only_drawing_column_does_not_become_spec_under_manual_mapping(self):
        source=self.paste('箱号\t物资名称\t零件图号\t数量\t单位\n1\t主轴护盖\t2S19252\t1\t个')
        source['mapping']={'spec':2}
        result=app.preview_post(self.request([source]))
        self.assertEqual(result['errors'],[])
        self.assertEqual((result['source_rows'][0]['spec'],result['source_rows'][0]['drawing']),('','2S19252'))
        self.assertEqual(result['rows'][0]['spec'],'')
        self.assertIn('图号：2S19252',result['rows'][0]['attribute'])

    def test_unrelated_page_box_conflict_still_requires_review(self):
        p=self.request([self.paste('装箱清单\n发运箱号：325-47\n箱号\t名称\t数量\t单位\n326\t扇形片\t100\t个')])
        r=app.preview_post(p)
        self.assertTrue(r['errors'])
        self.assertEqual(r['source_rows'][0]['box'],'326')
        self.assertEqual(r['source_rows'][0]['suggestions']['box'],'325-47')

    def test_arrival_only_needs_explicit_confirmation_and_box(self):
        p=self.request([self.paste('到货单号\t物资名称\t数量\t单位\nDH02\t垫片\t3\t件','arrival_note')])
        r=app.preview_post(p)
        self.assertTrue(r['errors'])
        p['inputs'][0].update(confirm_detail=True,defaults=dict(box='B03'))
        self.assertEqual(app.preview_post(p)['errors'],[])

    def test_formula_correction_and_preview_identity(self):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='装箱单'
        ws.append(['箱号','名称','数量','单位']);ws.append(['B04','测试物资','=2+3','件'])
        mem=io.BytesIO();wb.save(mem)
        p=self.request([dict(source='upload',role='packing_detail',filename='单据.xlsx',content=base64.b64encode(mem.getvalue()).decode())])
        r=app.preview_post(p)
        self.assertTrue(any('公式' in x for x in r['errors']))
        p['corrections']={r['source_rows'][0]['id']:dict(qty='5')}
        corrected=app.preview_post(p)
        self.assertFalse(corrected['errors'])
        p['recognition_id']=r['recognition_id']
        with self.assertRaisesRegex(ValueError,'不匹配'): app.import_post(p)
        p['recognition_id']=corrected['recognition_id']
        app.import_post(p)
        self.assertEqual(app.overview()['stock'][0]['qty'],5)

    def test_template_and_demo_are_separate(self):
        for name in ('到货单与装箱单模板.xlsx','演示数据-虚构勿入账.xlsx'):
            wb=openpyxl.load_workbook(app.ROOT/'static'/name)
            self.assertEqual(wb.sheetnames,['到货单','装箱单'])
            self.assertEqual([c.value for c in wb['到货单'][1]][:4],['到货单号','物资名称','数量','单位'])
            self.assertEqual([c.value for c in wb['装箱单'][1]][:5],['到货单号','箱号','物资名称','数量','单位'])
            wb.close()
        from tools.prepare_inbound_demo import demo_directory
        self.assertEqual(demo_directory(app.ROOT),app.ROOT/'demo-data')
        self.assertNotEqual(demo_directory(app.ROOT),app.ROOT/'data')

    def test_demo_seed_is_fictional_and_existing_demo_is_preserved(self):
        from tools.prepare_inbound_demo import prepare
        old_env={k:os.environ.get(k) for k in ('MATERIALS_DATA','MATERIALS_PORT','MATERIALS_DEMO_MODE')}
        try:
            with tempfile.TemporaryDirectory() as tmp:
                demo=prepare(Path(tmp))
                first=demo.overview()
                self.assertEqual(sum(r['qty'] for r in first['stock']),228)
                self.assertTrue(all(r['code'].startswith('DEMO-ONLY-') for r in first['stock']))
                self.assertEqual(len(first['documents']),1)
                prepare(Path(tmp))
                self.assertEqual(len(demo.overview()['documents']),1)
        finally:
            app.DATA=Path(self.tmp.name)
            for key,value in old_env.items():
                if value is None: os.environ.pop(key,None)
                else: os.environ[key]=value

    def test_demo_workbook_discrepancy_and_production_rejection(self):
        content=base64.b64encode((app.ROOT/'static'/'演示数据-虚构勿入账.xlsx').read_bytes()).decode()
        p=self.request([dict(source='upload',role=role,filename='演示数据.xlsx',content=content,sheet=sheet) for role,sheet in [('arrival_note','到货单'),('packing_detail','装箱单')]])
        old=os.environ.pop('MATERIALS_DEMO_MODE',None)
        try:
            self.assertTrue(any('演示' in e for e in app.preview_post(p)['errors']))
            os.environ['MATERIALS_DEMO_MODE']='1'
            r=app.preview_post(p)
            self.assertEqual(sum(not x['ok'] for x in r['reconciliation']),1)
            self.assertTrue(any('数量不一致' in e for e in r['errors']))
        finally:
            os.environ.pop('MATERIALS_DEMO_MODE',None)
            if old is not None: os.environ['MATERIALS_DEMO_MODE']=old

    def test_blocked_archive_can_reopen_saved_inputs_for_manual_review(self):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='装箱明细'
        ws.append(['箱号','物资名称','数量','单位']);ws.append(['','螺栓','4','件'])
        mem=io.BytesIO();wb.save(mem)
        content=base64.b64encode(mem.getvalue()).decode()
        source=dict(source='upload',role='packing_detail',filename='批次B01.xlsx',content=content,sheet='装箱明细')
        p=self.request([source],arrival_date='2026-09-30')
        p['original_name']='批次B01.xlsx'
        p['defaults']['batch']='B01'
        blocked=app.preview_post(p)
        self.assertTrue(blocked['errors'])

        resumed=app.recognition_reopen(blocked['recognition_id'])

        self.assertEqual(resumed['inputs'][0]['filename'],'批次B01.xlsx')
        self.assertEqual(resumed['inputs'][0]['sheet'],'装箱明细')
        self.assertEqual(resumed['defaults']['package'],'A1')
        self.assertEqual(resumed['defaults']['batch'],'B01')
        self.assertEqual(resumed['arrival_date'],'2026-09-30')
        self.assertEqual(resumed['original_name'],'批次B01.xlsx')
        self.assertEqual(resumed['inputs'][0]['content'],content)
        self.assertNotIn('mapping',resumed['inputs'][0])
        listed=app.recognition_list({})['rows'][0]
        self.assertEqual(listed['original_name'],'批次B01.xlsx')
        self.assertTrue(listed['can_reopen'])
        self.assertEqual(listed['issue_count'],len(blocked['errors']))
        self.assertIn('第2行',listed['issue_summary'])
        self.assertEqual(app.overview()['stock'],[])

    def test_archive_batch_delete_is_atomic_and_preserves_stock(self):
        app.admin_password_set({'new_password':'123'})
        p=self.request([self.paste('箱号\t物资名称\t数量\t单位\nB1\t螺栓\t2\t件')])
        app.import_post(p)
        posted=app.recognition_list({})['rows'][0]['id']
        pending=app.preview_post(self.request([self.paste('箱号\t物资名称\t数量\t单位\nB2\t钢板\t待核\t件')]))['recognition_id']
        before=app.overview()
        for payload in ({'ids':[posted,pending],'password':'wrong'},{'ids':[],'password':'123'},{'ids':[posted,999999],'password':'123'},{'ids':[posted,True],'password':'123'}):
            with self.assertRaises(ValueError): app.delete_recognitions(payload)
            self.assertEqual(app.recognition_list({})['total'],2)
            with app.connect() as c: self.assertEqual(c.execute('SELECT COUNT(*) FROM admin_actions').fetchone()[0],0)
        result=app.delete_recognitions({'ids':[posted,str(pending),posted],'password':'123'})
        self.assertEqual(result['deleted'],[posted,pending])
        self.assertEqual(app.recognition_list({})['total'],0)
        self.assertEqual(app.overview()['documents'],before['documents'])
        self.assertEqual(app.overview()['stock'],before['stock'])
        with app.connect() as c: self.assertEqual(c.execute('SELECT COUNT(*) FROM admin_actions').fetchone()[0],2)

    def test_archive_delete_password_preserves_receipt_and_shared_source(self):
        app.admin_password_set({'new_password':'123'})
        for status in ('READY','BLOCKED','CANCELLED'):
            p=self.request([self.paste(f'箱号\t物资名称\t数量\t单位\nB1\t{status}\t'+('1' if status=='READY' else '待核')+'\t件')])
            preview=app.preview_post(p);rec_id=preview['recognition_id']
            if status=='CANCELLED': app.cancel_recognition({'id':rec_id})
            record=app.recognition_detail(rec_id);source=app.DATA/record['source_path']
            original=source.read_bytes()
            self.assertEqual(record['status'],status)
            with self.assertRaisesRegex(ValueError,'密码错误'):
                app.delete_recognition({'id':rec_id,'password':'wrong'})
            self.assertEqual(app.recognition_detail(rec_id)['status'],status)
            app.delete_recognition({'id':rec_id,'password':'123'})
            with self.assertRaisesRegex(ValueError,'不存在'): app.recognition_detail(rec_id)
            self.assertEqual(source.read_bytes(),original)
        p=self.request([self.paste('箱号\t物资名称\t数量\t单位\nB2\t螺栓\t4\t件')])
        doc=app.import_post(p)
        record=app.recognition_detail(app.recognition_list({})['rows'][0]['id'])
        with app.connect() as c:
            copy=dict(c.execute('SELECT * FROM recognitions WHERE id=?',(record['id'],)).fetchone())
            copy.pop('id');copy['selection_hash']+='-copy'
            copy_id=c.execute('INSERT INTO recognitions('+','.join(copy)+') VALUES('+','.join('?' for _ in copy)+')',tuple(copy.values())).lastrowid
        before=app.overview();source=app.DATA/record['source_path'];original=source.read_bytes()
        app.delete_recognition({'id':record['id'],'password':'123'})
        self.assertEqual(app.overview()['documents'],before['documents'])
        self.assertEqual(app.overview()['stock'],before['stock'])
        self.assertEqual(app.recognition_detail(copy_id)['document_id'],doc['id'])
        with self.assertRaisesRegex(ValueError,'已导入'):
            app.import_post(dict(p,request_key=str(uuid.uuid4())))
        stock=app.overview()['stock'][0]
        issue=app.stock_post(dict(operator='TEST',kind='OUT',party='TEST CREW',purpose='Synthetic check',request_key=str(uuid.uuid4()),lines=[{'stock_id':stock['id'],'qty':'1'}]))
        with self.assertRaisesRegex(ValueError,'后续库存'):
            app.admin_delete({'id':doc['id'],'password':'123'})
        self.assertEqual(app.recognition_detail(copy_id)['status'],'POSTED')
        app.admin_delete({'id':issue['id'],'password':'123'})
        app.admin_delete({'id':doc['id'],'password':'123'})
        remaining=app.recognition_detail(copy_id)
        self.assertEqual((remaining['status'],remaining['document_id']),('CANCELLED',None))
        self.assertFalse(app.overview()['documents'])
        self.assertFalse(app.overview()['stock'])
        self.assertEqual(source.read_bytes(),original)
        with app.connect() as c:
            audits=c.execute("SELECT reason FROM admin_actions WHERE action='RECOGNITION_DELETE'").fetchall()
            self.assertEqual(len(audits),4)
            self.assertEqual(json.loads(audits[-1]['reason'])['id'],record['id'])


if __name__ == '__main__':
    unittest.main()
