"""Local materials ledger. Exact quantities, transactional posting, immutable documents."""
import base64, csv, gzip, hashlib, hmac, io, json, mimetypes, os, re, secrets, sqlite3, sys, threading, unicodedata, zipfile
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote, quote
import openpyxl
import local_ai

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('MATERIALS_DATA', ROOT / 'data'))
SCALE = 1000000
STATES = ('AVAILABLE', 'PENDING', 'QUARANTINE', 'SPARE')
TRANSPORT = ('PLANNED', 'DISPATCHED', 'SEA', 'CUSTOMS', 'RELEASED', 'ROAD', 'ARRIVED')
TOKEN = secrets.token_urlsafe(32)
APP_VERSION = '0.1.36'  # keep in sync with CHANGELOG.md and static/app.js
RECOGNITION_PARSER_VERSION = '1'
RECOGNITION_TIMEZONE = timezone(timedelta(hours=2), 'SAST')
# The ledger is the only unbounded table: one row per movement.  Sending all of
# it on every refresh is what made large databases slow to open, so the default
# state ships only the most recent slice and the rest is paged in on demand.
LEDGER_PAGE_SIZE = 300
LEDGER_MAX_PAGE_SIZE = 2000
# The item master is the other table that grows with the catalogue, and it had
# the same problem twice over: every refresh shipped every item row, and every
# row travelled a second time nested inside its contract group.  At 3,000 items
# the two copies were 2.07 MB and effectively the whole response.  Contracts now
# carry codes plus a count, and item rows travel as a bounded page that the
# materials table pages in on demand - the same shape as the ledger.
ITEMS_PAGE_SIZE = 300
ITEMS_MAX_PAGE_SIZE = 2000
# Columns the server-side item search looks at.  Keep in sync with the search
# placeholder on the materials table.
ITEM_SEARCH_FIELDS = ('code', 'name', 'spec', 'unit', 'package', 'attribute')

# The project uses a fixed set of equipment packages/sections.  Keep the
# short A-code stable for filtering and documents, while exposing the formal
# Chinese and English names to both local and international users.
SEGMENTS = (
    ('A1', '光伏组件及组件配套材料', 'PV Modules and Module Accessories'),
    ('A2', '单轴跟踪支架及跟踪控制系统', 'Single-Axis Tracker Mounting System and Tracker Control System'),
    ('A3', '光伏组串式逆变器', 'PV String Inverters'),
    ('A4', 'PV箱式变压器及箱变成套设备', 'PV Box-Type Transformers and Packaged Substation Equipment'),
    ('A5', 'BESS储能区设备', 'BESS Area Equipment'),
    ('A6', '高低压配电及动力控制设备', 'HV/LV Power Distribution and Power Control Equipment'),
    ('A7', '综合自动化系统', 'Integrated Automation System'),
    ('A8', '光伏及BESS区域辅助变压器', 'Auxiliary Transformers for PV and BESS Areas'),
    ('A9', '直流及UPS系统', 'DC and UPS Systems'),
    ('A10', '光伏专用缆', 'PV Cables'),
    ('A11', '全站电力、控制、通信光缆及附件', 'Plant-Wide Power, Control and Fiber-Optic Communication Cables and Accessories'),
    ('A12', 'HPC/PPC、AGC/AVC及并网控制系统', 'HPC/PPC, AGC/AVC and Grid-Connection Control Systems'),
    ('A13', '全站SCADA及数据管理系统', 'Plant-Wide SCADA and Data Management System'),
    ('A14', '气象站', 'Meteorological Station'),
    ('A15', '132kV主变压器', '132 kV Main Transformers'),
    ('A16', '无功补偿装置', 'Reactive Power Compensation Equipment'),
)
SEGMENT_CODES = {code for code, _, _ in SEGMENTS}
SEGMENT_NAMES = {code: {'zh': zh, 'en': en} for code, zh, en in SEGMENTS}

# Explicit warehouse A1-A16 to current procurement Lot 01-16 crosswalk.
LOT_CROSSWALK = {
    'A1':'Lot-01', 'A2':'Lot-02', 'A3':'Lot-03', 'A4':'Lot-04',
    'A5':'Lot-05', 'A6':'Lot-06', 'A7':'Lot-07', 'A8':'Lot-08',
    'A9':'Lot-09', 'A10':'Lot-10', 'A11':'Lot-11', 'A12':'Lot-12',
    'A13':'Lot-13', 'A14':'Lot-14', 'A15':'Lot-15', 'A16':'Lot-16',
}
PROJECT_ARRIVAL_TARGETS = (
    {'view':'overview','key':'bess-cabinet','name':'储能柜','package':'A5','unit':'台',
     'aliases':('储能柜','电池柜','储能电池柜','电池舱'),'model_display_count':86},
    {'view':'overview','key':'pcs','name':'PCS 储能变流器','package':'A5','unit':'台',
     'aliases':('PCS','储能PCS','储能变流器','储能变流升压一体机')},
    {'view':'pv','key':'pv-inverter','name':'组串式逆变器','package':'A3','unit':'台',
     'aliases':('光伏组串式逆变器','组串式逆变器','光伏逆变器','逆变器')},
    {'view':'pv','key':'pv-transformer','name':'PV 箱式变压器','package':'A4','unit':'台',
     'aliases':('PV箱式变压器','PV箱变','光伏箱变','箱式变压器','箱变')},
    {'view':'transmission','key':'main-transformer','name':'132 kV 主变压器','package':'A15','unit':'台',
     'aliases':('132kV主变压器','132kV主变','主变压器','主变')},
)

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def clean(v): return str(v if v is not None else '').strip()
def required(v, label, max_len=250):
    s = clean(v)
    if not s or len(s) > max_len: raise ValueError('必填或过长 / Required or too long: ' + label)
    return s
def units(v, positive=True):
    try: d = Decimal(clean(v))
    except InvalidOperation: raise ValueError('数量格式错误 / Invalid quantity')
    if not d.is_finite() or (positive and d <= 0) or abs(d) > Decimal('1000000000') or d*SCALE != (d*SCALE).to_integral_value():
        raise ValueError('数量须为正数且最多六位小数 / Positive quantity, maximum six decimal places')
    return int(d*SCALE)
def number(v): return float(Decimal(v)/SCALE)
def parse_box_contents(text):
    """Read simple one-item-per-line contents from a box remark.

    These are child details for display and traceability. They do not create
    independent stock balances; the parent carton line remains the quantity
    that is posted to inventory.
    """
    result=[]
    for raw in clean(text).splitlines():
        line=raw.strip().strip('（）()[]【】,，;；')
        if not line: continue
        m=re.match(r'^(.*?)(\d+(?:\.\d+)?)\s*(件|个|套|箱|台|根|米|块|卷|包|只|把|组|张|支|颗|枚|片|对|批|吨|kg|KG|m|M)\s*[）)]?$',line,re.I)
        if not m: continue
        desc=clean(m.group(1)).strip('：: -')
        if not desc: continue
        code_match=re.search(r'(?<![A-Za-z0-9])([A-Z]{2,}[A-Z0-9]*(?:-[A-Z0-9]+)+)',desc)
        code=code_match.group(1) if code_match else ''
        name=clean(desc.replace(code,' ')).strip(' -') if code else desc
        result.append({'name':name,'qty':m.group(2),'unit':m.group(3),'code':code,'raw':line})
    return result
def legacy_device_code(raw):
    s=clean(raw)
    aliases={'光伏组件':'组件','组件':'组件','支架':'支架','电池':'电池','电池舱':'电池舱','储能':'储能','PCS':'PCS','箱变':'箱变','变压器':'变压器','汇流箱':'汇流箱','逆变器':'逆变器','电缆':'电缆','轨道机器人':'轨道机器人','轮式机器人':'轮式机器人','无人机':'无人机','无人机机场':'无人机机场','水浸传感器':'水浸传感器','摄像头':'摄像头'}
    label=aliases.get(s,s)
    label=re.sub(r'[\\/:*?"<>|]+','',label)
    label=re.sub(r'\s+','',label)
    return label[:40] or '设备'

def generated_batch(p):
    # New UI batches are deliberately short (B01, B02...).  Keep the old
    # device/date form only for legacy direct API callers that still send the
    # deprecated batch_device field; no new receipt can create that form.
    value=p.get('batch_number',p.get('batchNumber','1'))
    seq=clean(value)
    if isinstance(value,bool) or not re.fullmatch(r'[0-9]{1,12}',seq) or int(seq)<=0:
        raise ValueError('批次序号须为正整数（最多12位） / Batch sequence must be a positive integer (maximum 12 digits)')
    seq=f'{int(seq):02d}'
    legacy_device=clean(p.get('batch_device') or p.get('batchDevice') or p.get('device_type'))
    if legacy_device and not p.get('segment_required'):
        raw_date=clean(p.get('arrival_date') or p.get('batch_date'))
        digits=''.join(ch for ch in raw_date if ch.isdigit())
        yymmdd=(digits[2:8] if len(digits)>=8 else datetime.now().strftime('%y%m%d'))
        return f"{legacy_device_code(legacy_device)}-{yymmdd}-{seq}"
    return f'B{seq}'
class ClosingConnection(sqlite3.Connection):
    def __exit__(self, *args):
        try: return super().__exit__(*args)
        finally: self.close()

def connect():
    DATA.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DATA/'materials.db', timeout=20, factory=ClosingConnection)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    c.execute('PRAGMA journal_mode=WAL')
    return c
def init():
    with connect() as c:
        c.executescript('''
CREATE TABLE IF NOT EXISTS items(code TEXT PRIMARY KEY, name TEXT NOT NULL, spec TEXT NOT NULL, unit TEXT NOT NULL, package TEXT NOT NULL, attribute TEXT NOT NULL DEFAULT '');
CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY, number TEXT UNIQUE NOT NULL, kind TEXT NOT NULL, external_ref TEXT NOT NULL, operator TEXT NOT NULL, party TEXT NOT NULL, purpose TEXT NOT NULL, created TEXT NOT NULL, arrival_date TEXT NOT NULL DEFAULT '', request_key TEXT UNIQUE NOT NULL, file_hash TEXT, source_path TEXT, reversed_by INTEGER, reverse_of INTEGER UNIQUE);
CREATE TABLE IF NOT EXISTS recognitions(id INTEGER PRIMARY KEY, recognized_at TEXT NOT NULL, updated_at TEXT NOT NULL, business_date TEXT NOT NULL DEFAULT '', category TEXT NOT NULL, status TEXT NOT NULL, original_name TEXT NOT NULL, raw_sha256 TEXT NOT NULL, selection_hash TEXT NOT NULL, sheet_name TEXT NOT NULL, parser_version TEXT NOT NULL, source_path TEXT NOT NULL, segments_json TEXT NOT NULL DEFAULT '[]', batches_json TEXT NOT NULL DEFAULT '[]', snapshot_json TEXT NOT NULL, document_id INTEGER);
CREATE UNIQUE INDEX IF NOT EXISTS unique_recognition_version ON recognitions(raw_sha256,selection_hash,parser_version);
CREATE INDEX IF NOT EXISTS idx_recognitions_date ON recognitions(recognized_at DESC,id DESC);
CREATE INDEX IF NOT EXISTS idx_recognitions_status ON recognitions(status,category,recognized_at DESC);
CREATE UNIQUE INDEX IF NOT EXISTS unique_import_file ON documents(file_hash) WHERE file_hash IS NOT NULL;
CREATE UNIQUE INDEX IF NOT EXISTS unique_business_ref ON documents(kind,external_ref) WHERE kind IN ('IN','BASELINE');
CREATE TABLE IF NOT EXISTS lines(id INTEGER PRIMARY KEY, doc_id INTEGER NOT NULL REFERENCES documents(id), code TEXT NOT NULL, name TEXT NOT NULL, spec TEXT NOT NULL, unit TEXT NOT NULL, batch TEXT NOT NULL, box TEXT NOT NULL DEFAULT '', warehouse TEXT NOT NULL, bin TEXT NOT NULL, state TEXT NOT NULL, qty INTEGER NOT NULL, attribute TEXT NOT NULL DEFAULT '', contents TEXT NOT NULL DEFAULT '', source_row INTEGER);
CREATE TABLE IF NOT EXISTS stock(id INTEGER PRIMARY KEY, code TEXT NOT NULL REFERENCES items(code), batch TEXT NOT NULL, box TEXT NOT NULL DEFAULT '', warehouse TEXT NOT NULL, bin TEXT NOT NULL, state TEXT NOT NULL, qty INTEGER NOT NULL CHECK(qty>=0), contents TEXT NOT NULL DEFAULT '', UNIQUE(code,batch,box,warehouse,bin,state));
CREATE TABLE IF NOT EXISTS movements(id INTEGER PRIMARY KEY, doc_id INTEGER NOT NULL REFERENCES documents(id), stock_id INTEGER NOT NULL REFERENCES stock(id), delta INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS design(code TEXT PRIMARY KEY REFERENCES items(code), qty INTEGER NOT NULL CHECK(qty>=0));
CREATE TABLE IF NOT EXISTS changes(id INTEGER PRIMARY KEY, code TEXT NOT NULL REFERENCES design(code), delta INTEGER NOT NULL, reason TEXT NOT NULL, operator TEXT NOT NULL, created TEXT NOT NULL, request_key TEXT UNIQUE NOT NULL);
CREATE TABLE IF NOT EXISTS batches(code TEXT PRIMARY KEY, supplier TEXT NOT NULL, status TEXT NOT NULL, updated TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, batch TEXT NOT NULL, old_status TEXT, new_status TEXT, operator TEXT NOT NULL, note TEXT NOT NULL, created TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS admin_actions(id INTEGER PRIMARY KEY, action TEXT NOT NULL, document_id INTEGER NOT NULL, document_number TEXT NOT NULL, reason TEXT NOT NULL, operator TEXT NOT NULL, created TEXT NOT NULL);

-- Business indexes.  Without these, every per-item summary falls back to a
-- full scan of lines/movements, so refresh time grows with (items x rows).
-- IF NOT EXISTS also back-fills indexes on databases created earlier.
CREATE INDEX IF NOT EXISTS idx_lines_code ON lines(code);
CREATE INDEX IF NOT EXISTS idx_lines_doc ON lines(doc_id);
CREATE INDEX IF NOT EXISTS idx_movements_stock ON movements(stock_id);
CREATE INDEX IF NOT EXISTS idx_movements_doc ON movements(doc_id);
CREATE INDEX IF NOT EXISTS idx_changes_code ON changes(code);
''')
        for table,column,definition in [('items','attribute','TEXT NOT NULL DEFAULT \'\''),('documents','arrival_date','TEXT NOT NULL DEFAULT \'\''),('lines','box','TEXT NOT NULL DEFAULT \'\''),('lines','attribute','TEXT NOT NULL DEFAULT \'\''),('lines','contents','TEXT NOT NULL DEFAULT \'\''),('stock','box','TEXT NOT NULL DEFAULT \'\''),('stock','contents','TEXT NOT NULL DEFAULT \'\'')]:
            cols={row['name'] for row in c.execute(f'PRAGMA table_info({table})').fetchall()}
            if column not in cols: c.execute(f'ALTER TABLE {table} ADD COLUMN {column} {definition}')
        # Older trial databases used a stock key without the box number.  A
        # simple ALTER cannot change that UNIQUE constraint, so rebuild the
        # table in place while preserving existing balances and foreign-key
        # names.  This keeps two boxes of the same item as separate positions
        # after an upgrade instead of silently merging them.
        stock_sql=c.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='stock'").fetchone()['sql'] or ''
        compact=''.join(stock_sql.lower().split())
        if 'unique(code,batch,box,warehouse,bin,state)' not in compact:
            c.commit()
            c.execute('PRAGMA foreign_keys=OFF')
            c.execute('''CREATE TABLE stock_new(id INTEGER PRIMARY KEY, code TEXT NOT NULL REFERENCES items(code), batch TEXT NOT NULL, box TEXT NOT NULL DEFAULT '', warehouse TEXT NOT NULL, bin TEXT NOT NULL, state TEXT NOT NULL, qty INTEGER NOT NULL CHECK(qty>=0), contents TEXT NOT NULL DEFAULT '', UNIQUE(code,batch,box,warehouse,bin,state))''')
            c.execute('''INSERT INTO stock_new(id,code,batch,box,warehouse,bin,state,qty,contents)
                         SELECT id,code,batch,COALESCE(box,''),warehouse,bin,state,qty,COALESCE(contents,'') FROM stock''')
            c.execute('DROP TABLE stock')
            c.execute('ALTER TABLE stock_new RENAME TO stock')
            c.execute('PRAGMA foreign_keys=ON')
        # Upgrade earlier trial imports: turn preserved long remarks into
        # structured child details without changing the parent stock quantity.
        for row in c.execute("SELECT s.id,s.contents,i.attribute FROM stock s JOIN items i ON i.code=s.code WHERE COALESCE(i.attribute,'')<>''").fetchall():
            child=parse_box_contents(row['attribute'])
            if child: c.execute('UPDATE stock SET contents=? WHERE id=?',(json.dumps(child,ensure_ascii=False),row['id']))
        for row in c.execute("SELECT l.id,l.contents,l.attribute FROM lines l WHERE COALESCE(l.attribute,'')<>''").fetchall():
            child=parse_box_contents(row['attribute'])
            if child: c.execute('UPDATE lines SET contents=? WHERE id=?',(json.dumps(child,ensure_ascii=False),row['id']))
        # A previous trial could leave item-master rows behind after its only
        # receipt was removed.  Remove only truly orphaned zero-balance items;
        # design baselines, changes, other lines and non-zero stock preserve
        # their item records.
        for row in c.execute('SELECT code FROM items').fetchall():
            cleanup_orphan_item(c,row['code'])

ALIASES = {
'code':['物资编码','材料编码','设备编码','物料编码','编码','料号','code','itemcode','materialcode','sku','partnumber'],
 'name':['物资名称','材料名称','设备名称','品名','名称','零件名称','货物名称','name','description','itemname','product'],
 'spec':['规格型号','规格','型号','零件图号','图号','spec','specification','model'],
 'unit':['单位','计量单位','unit','uom'],
 'qty':['数量','入库数量','实收数量','设计数量','总量','quantity','qty','receivedquantity'],
 'warehouse':['仓库','库房','warehouse'], 'bin':['库位','货位','位置','bin','location','binlocation'],
 'batch':['批次','发运批次','批次号','合同批次','batch','lot','lotnumber'], 'box':['箱号','箱编号','箱件号','箱号/箱件号','装箱编号','装箱单号','箱单号','box','carton','cartonno'], 'package':['合同','合同号','合同编号','合同名称','设备包','设计项','标段','项目号','项目名称','package'], 'attribute':['设备属性','厂家属性','装配描述','装配说明','备注','说明','attribute','attributes','remarks','notes'],
 'state':['质量状态','库存状态','state','status']}
def norm(v): return ''.join(ch for ch in clean(v).lower() if ch.isalnum())
def contract_numbers(rows):
    found=[]
    for row in rows:
        for value in row:
            text=clean(value)
            if not text: continue
            match=re.search(r'(?:合同编号|合同号|合同)\s*[:：]?\s*([A-Za-z0-9][A-Za-z0-9._/-]*)',text)
            if match and match.group(1) not in found: found.append(match.group(1))
    return found

def parse_file(p):
    filename = required(p.get('filename'), '文件名 / Filename')
    if Path(filename).name.startswith('~$'): raise ValueError('这是 Excel 锁文件，请上传实际工作簿')
    try: raw = base64.b64decode(p.get('content',''), validate=True)
    except Exception: raise ValueError('文件编码无效 / Invalid file encoding')
    if not raw or len(raw)>12*1024*1024: raise ValueError('文件为空或超过12MB / Empty file or exceeds 12 MB')
    ext=Path(filename).suffix.lower()
    if ext == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>80*1024*1024: raise ValueError('解压文件过大 / Expanded workbook too large')
            # Normal mode is intentional here: manufacturer workbooks often use
            # merged cells for headings and section context. We preserve the
            # source row and never evaluate formulas.
            wb=openpyxl.load_workbook(io.BytesIO(raw), read_only=False, data_only=False)
            sheets=wb.sheetnames
            sheet=p.get('sheet') or sheets[0]
            if sheet not in sheets: raise ValueError('工作表不存在 / Sheet not found')
            # A workbook can contain both a carton-level shipment list and
            # one or more packing-detail sheets.  The shipment list carries
            # carton quantities and must never be posted as receipt lines;
            # receipt posting is intentionally limited to the packing-detail
            # sheet selected by the operator.
            ws=wb[sheet]
            # Contract number is often printed in the cover/header area rather
            # than repeated in every packing-detail row. Read the small cover
            # area of every worksheet so a workbook with一期/二期 sheets can
            # still provide one unambiguous contract default.
            workbook_meta=[]
            for meta_ws in wb.worksheets:
                for row in meta_ws.iter_rows(min_row=1,max_row=min(meta_ws.max_row,25),max_col=min(meta_ws.max_column,40),values_only=True):
                    workbook_meta.append(list(row))
            if ws.max_row and ws.max_row>10002: raise ValueError('最多10000条明细 / Maximum 10000 rows')
            rows=[]
            for row in ws.iter_rows():
                if len(rows)>10001: raise ValueError('行数过多 / Too many rows')
                rows.append([{'formula':True} if cell.data_type=='f' else clean(cell.value) for cell in row[:80]])
            wb.close()
        except ValueError: raise
        except Exception: raise ValueError('无法读取xlsx / Cannot read workbook')
    elif ext=='.csv':
        try: txt=raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: txt=raw.decode('gb18030')
            except UnicodeDecodeError: raise ValueError('请另存为UTF-8 CSV / Please save as UTF-8 CSV')
        rows=list(csv.reader(io.StringIO(txt))); sheets=['CSV']; sheet='CSV'; workbook_meta=rows[:25]
        if len(rows)>10002: raise ValueError('最多10000条明细 / Maximum 10000 rows')
    else: raise ValueError('支持xlsx和csv；xls请另存 / Use xlsx or csv; convert legacy xls')
    requested=int(p.get('header',0) or 0)
    if requested<0 or requested>len(rows): raise ValueError('表头行无效 / Invalid header row')
    shipment_tokens=('发货','发运','shipment','carton')
    packing_tokens=('装箱','箱单','packing')
    worksheet_shipment_hint=any(token in sheet.lower() for token in shipment_tokens)
    worksheet_packing_hint=any(token in sheet.lower() for token in packing_tokens)
    filename_shipment_hint=any(token in filename.lower() for token in shipment_tokens)
    heading_text=''.join(norm(value) for row in rows[:8] for value in row if not isinstance(value,dict))
    heading_shipment_hint=any(token in heading_text for token in ('发货清单','发运单','shipmentlist','shippinglist'))
    heading_packing_hint=any(token in heading_text for token in ('装箱清单','装箱明细','箱单','packinglist'))
    alias_norms={key:{norm(a) for a in aliases} for key,aliases in ALIASES.items()}
    def infer(row):
        out={}
        for key in ALIASES:
            hits=[i for i,h in enumerate(row) if norm(h) in alias_norms[key]]
            # Attributes may legitimately have both an assembly-description
            # column and a remarks column; prefer the descriptive column.
            out[key]=hits[0] if (len(hits)==1 or (key=='attribute' and hits)) else -1
        return out
    # Factory workbooks repeat the same table header on every packing page.
    # Header 0 means automatic detection; an entered header is still honored
    # when it points at a valid table header.
    candidates=[i for i,row in enumerate(rows) if all(infer(row)[k]>=0 for k in ('name','qty','unit'))]
    if requested:
        chosen=[requested-1] if requested-1 in candidates else candidates
    else:
        chosen=candidates
    if not chosen:
        shipment_hint=worksheet_shipment_hint or heading_shipment_hint or (filename_shipment_hint and not (worksheet_packing_hint or heading_packing_hint))
        if shipment_hint and rows:
            chosen=[requested-1] if requested and requested-1<len(rows) else [next((i for i,row in enumerate(rows) if any(clean(x) for x in row)),0)]
        if requested and requested-1<len(rows): chosen=[requested-1]
        elif not chosen: raise ValueError('未找到包含名称、数量、单位的表头 / Could not find a header with name, quantity and unit')
    headers=[clean(x) for x in rows[chosen[0]]]
    header_norms={norm(h) for h in headers if clean(h)}
    # Shipment lists describe cartons and transport parameters.  Packing
    # sheets describe the item rows inside a carton.  Use strong header
    # markers in addition to the worksheet/file name so a generically named
    # worksheet is still routed safely.
    shipment_markers={norm(x) for x in ('包装类型','箱体尺寸','箱尺寸','外形尺寸','长宽高','毛重','净重','箱数','装箱数','包装参数')}
    packing_detail_hint=any(norm(h) in {norm('物料编码'),norm('物资编码'),norm('零件名称'),norm('零件图号')} for h in headers)
    shipment_hint=worksheet_shipment_hint or heading_shipment_hint or (filename_shipment_hint and not (worksheet_packing_hint or heading_packing_hint or packing_detail_hint))
    sheet_kind='shipment' if shipment_hint or any(marker in header_norms for marker in shipment_markers) else 'packing'
    inferred=infer(rows[chosen[0]])
    mapping=p.get('mapping') if p.get('mapping') is not None else inferred
    defaults=dict(p.get('defaults',{}) or {})
    detected_contracts=contract_numbers(workbook_meta if ext=='.xlsx' else rows)
    # A selected fixed segment is authoritative for this receipt.  For legacy
    # API callers that do not select a segment, preserve the old cover-number
    # fallback so existing evidence can still be previewed and tested.
    if not clean(defaults.get('package')) and len(detected_contracts)==1 and not p.get('segment_required'):
        defaults['package']=detected_contracts[0]
    if p.get('kind','IN')=='IN':
        # Factory packing sheets normally contain only source facts.  Apply
        # one receipt-wide batch and location default so the administrator
        # does not have to repeat those fields on every row.
        defaults['batch']=clean(defaults.get('batch')) or generated_batch(p)
        defaults['warehouse']=clean(defaults.get('warehouse')) or '主仓库'
        defaults['bin']=clean(defaults.get('bin')) or '待分配'
        defaults['state']=clean(defaults.get('state')) or 'AVAILABLE'
    data=[]; errors=[]; warnings=[]; generated_codes=set(); reused_codes=set(); generated_by_identity={}; skipped_layout=0; merged_rows=0
    if sheet_kind=='shipment' and p.get('kind','IN') in ('IN','BASELINE'):
        errors.append('发货清单是箱级物流核对表，已归档但不能直接办理入库 / Shipment lists are archived for logistics review and cannot be posted as receipts')
    selected_segment=clean(defaults.get('package')).upper()
    if p.get('segment_required') and selected_segment not in SEGMENT_CODES:
        errors.append('请选择固定标段 A1-A16 / Select one fixed segment A1-A16')
    existing_identity={}
    with connect() as c:
        for old in c.execute('SELECT code,name,spec,unit,package,attribute FROM items').fetchall():
            identity=tuple(clean(old[k]) for k in ('name','spec','unit','package','attribute'))
            existing_identity.setdefault(identity,[]).append(old['code'])
    for sec,hrow in enumerate(chosen):
        section_mapping=mapping if p.get('mapping') is not None else infer(rows[hrow])
        # A factory box list may expose both “装配描述” and “备注”.  Keep
        # both columns when mapping is automatic so the long contents of a
        # carton remain available for receipt checking and later lookup.
        attribute_columns=[i for i,h in enumerate(rows[hrow]) if norm(h) in alias_norms['attribute']]
        selected_attribute=int(section_mapping.get('attribute',-1))
        if selected_attribute >= 0 and selected_attribute not in attribute_columns:
            attribute_columns.append(selected_attribute)
        end=chosen[sec+1] if sec+1<len(chosen) else len(rows)
        carry={'package':'','box':''}
        for idx in range(hrow+1,end):
            row=rows[idx]
            if not any(clean(x) for x in row): continue
            record={'source_row':idx+1}
            for key in ALIASES:
                col=int(section_mapping.get(key,-1))
                val=row[col] if 0<=col<len(row) else ''
                if isinstance(val,dict): errors.append(f'第{idx+1}行 / Row {idx+1}: 公式须转为值 / Convert formulas to values'); val=''
                # The operator-selected fixed segment applies to the whole
                # workbook.  Factory cover rows often contain a contract
                # number in the same column; that number must not overwrite
                # the selected A1-A16 segment.
                if key=='package' and clean(defaults.get('package')):
                    record[key]=clean(defaults.get('package'))
                else:
                    record[key]=clean(val) or clean(defaults.get(key,''))
            if attribute_columns:
                attrs=[]
                for col in attribute_columns:
                    val=row[col] if 0<=col<len(row) else ''
                    if isinstance(val,dict):
                        errors.append(f'第{idx+1}行 / Row {idx+1}: 公式须转为值 / Convert formulas to values')
                        continue
                    val=clean(val)
                    if val and val not in attrs: attrs.append(val)
                record['attribute']='；'.join(attrs) or clean(defaults.get('attribute',''))
            record['contents']=parse_box_contents(record.get('attribute',''))
            # Excel users often merge the contract or box cell across several
            # detail rows. openpyxl exposes the merged children as blank cells;
            # carry those two identifiers forward within the current page.
            for key in ('package','box'):
                if record[key]: carry[key]=record[key]
                elif carry[key]: record[key]=carry[key]
            # Packing sheets reserve serial-number-only rows and signature
            # rows inside each page. They are layout rows, not material lines.
            if not any(record[k] for k in ('code','name','qty')): continue
            if not record['name'] and not record['code']: continue
            try: amount=units(record['qty'])
            except (ValueError,TypeError) as e:
                # Footers and page metadata can occupy the table columns. A
                # material line is only actionable when its quantity is numeric;
                # retain a warning so the administrator can inspect it.
                skipped_layout += 1
                continue
            if not record['code'] and record['name']:
                # Many factory lists omit a material code. Give each distinct
                # item a readable code. If the same identity was already
                # received in an earlier batch, reuse its code so cross-batch
                # stock and issue searches still refer to one material.
                identity=tuple(record[k] for k in ('name','spec','unit','package','attribute'))
                basis='|'.join(identity)
                if basis not in generated_by_identity:
                    matches=existing_identity.get(identity,[])
                    if len(matches)==1:
                        generated_by_identity[basis]=matches[0]; reused_codes.add(matches[0])
                    else:
                        seq=len(generated_by_identity)+1
                        prefix=clean(defaults.get('batch')) or 'ITEM'
                        generated_by_identity[basis]=f'{prefix}-{seq:03d}'
                record['code']=generated_by_identity[basis]
                if record['code'] not in reused_codes: generated_codes.add(record['code'])
            record['state']=record['state'] or 'AVAILABLE'
            record['state']={'可用':'AVAILABLE','合格':'AVAILABLE','待检':'PENDING','隔离':'QUARANTINE','不合格':'QUARANTINE','备品':'SPARE','预留备品':'SPARE'}.get(record['state'],record['state'].upper())
            try:
                required(record['code'],'code')
                required(record['name'],'name',2000)
                required(record['unit'],'unit')
                record['amount']=amount
                if p.get('kind','IN')=='IN':
                    for key in ('batch','box','warehouse','bin'): required(record[key],key)
                    if record['state'] not in STATES: raise ValueError('库存状态无效 / Invalid stock state')
            except (ValueError,TypeError) as e: errors.append(f'第{idx+1}行 / Row {idx+1}: {e}')
            data.append(record)
    if generated_codes:
        sample=', '.join(sorted(generated_codes)[:5])
        more='…' if len(generated_codes)>5 else ''
        warnings.append(f'有 {len(generated_codes)} 个物资未提供物资编码，已按到货批次生成可追溯编号（示例 {sample}{more}）；请确认前核对 / {len(generated_codes)} items had no source item code; readable batch-sequence codes were generated. Verify before posting.')
    if reused_codes:
        warnings.append(f'有 {len(reused_codes)} 个无厂家编码物资沿用了既有内部编码，便于跨批次合并查询；请确认前核对 / {len(reused_codes)} code-less items reused existing internal codes for cross-batch grouping; verify before posting.')
    if skipped_layout:
        warnings.append(f'自动忽略 {skipped_layout} 行页眉、签字或分页内容 / Ignored {skipped_layout} header, signature or page-layout rows automatically.')
    if len(detected_contracts)==1 and not clean((p.get('defaults') or {}).get('package')):
        warnings.append(f'已从文件封面识别合同编号 {detected_contracts[0]}，并带入全部箱内明细 / Contract number {detected_contracts[0]} was read from the workbook cover and applied to the detail rows.')
    elif len(detected_contracts)>1 and not clean((p.get('defaults') or {}).get('package')):
        warnings.append('文件中识别到多个合同编号，请在入库前选择一个合同/标段 / Multiple contract numbers were found; select the contract/lot before posting.')
    if not data: errors.append('没有物资明细 / No material rows')
    combined={}
    for r in data:
        key=(r['code'],r['batch'],r['box'],r['warehouse'],r['bin'],r['state'])
        old=combined.get(key)
        if old:
            if any(old[k]!=r[k] for k in ('name','spec','unit')):
                errors.append(f"第{r['source_row']}行与第{old['source_row']}行编码信息冲突；请核对名称/规格/单位 / Conflicting item identity")
            old['amount'] += r['amount']; old['qty']=str(number(old['amount'])); merged_rows += 1
        else:
            combined[key]=r
    data=list(combined.values())
    if merged_rows:
        warnings.append(f'已将 {merged_rows} 行相同物资/批次/库位明细合计数量 / Aggregated {merged_rows} repeated item-position rows into their quantities.')
    # A workbook may legitimately contain both a packing-detail sheet and a
    # shipment/carton sheet. Treat each selected worksheet as its own import
    # while retaining the original workbook as the source evidence.
    digest=hashlib.sha256(raw+b'\0'+sheet.encode('utf-8')).hexdigest()
    with connect() as c:
        if c.execute('SELECT 1 FROM documents WHERE file_hash=?',(digest,)).fetchone(): errors.append('此文件已导入 / File already posted')
        for r in data:
            old=c.execute('SELECT * FROM items WHERE code=?',(r['code'],)).fetchone()
            if old and (old['unit']!=r['unit'] or old['spec']!=r['spec'] or old['name']!=r['name']): errors.append(f"{r['code']}: 名称/规格/单位与现有编码不符 / Item identity mismatch")
    return {'headers':headers,'sheets':sheets,'sheet':sheet,'sheet_kind':sheet_kind,'header_row':chosen[0]+1,'mapping':mapping,'rows':data,'errors':list(dict.fromkeys(errors))[:100],'warnings':list(dict.fromkeys(warnings))[:100],'count':len(data),'hash':digest},raw

def store_source(raw,filename):
    raw_hash=hashlib.sha256(raw).hexdigest()
    suffix=Path(filename).suffix.lower() or '.bin'
    rel='sources/'+raw_hash+suffix
    dest=DATA/rel
    if not dest.exists():
        dest.parent.mkdir(parents=True,exist_ok=True)
        tmp=dest.with_name(dest.name+'.'+secrets.token_hex(4)+'.tmp')
        tmp.write_bytes(raw)
        try: tmp.replace(dest)
        finally: tmp.unlink(missing_ok=True)
    return raw_hash,rel

def demo_mode():
    return os.environ.get('MATERIALS_DEMO_MODE')=='1'


def parse_intake(p):
    if 'inputs' in p:
        if p.get('kind','IN')!='IN': raise ValueError('粘贴和手动录入仅用于到货入库')
        import intake
        return intake.parse(p,sys.modules[__name__])
    result,raw=parse_file(p)
    return result,raw,p['filename']


def intake_ai_post(p):
    if p.get('kind')!='IN' or 'inputs' not in p: raise ValueError('本地AI仅辅助到货入库识别')
    result,_,_=parse_intake(p)
    return local_ai.suggest(result)


def recognition_category(p,result):
    if result.get('sheet_kind')=='shipment': return 'shipment'
    if p.get('kind')=='BASELINE': return 'baseline'
    if not result.get('rows') and result.get('errors'): return 'unclassified'
    if result.get('sheet_kind')=='packing': return 'packing'
    return 'unclassified'

def upsert_recognition(c,p,result,raw):
    raw_hash,source_path=store_source(raw,p.get('filename','source.bin'))
    parser_version=RECOGNITION_PARSER_VERSION
    row=c.execute('SELECT * FROM recognitions WHERE raw_sha256=? AND selection_hash=? AND parser_version=?',
                  (raw_hash,result.get('selection_hash',result['hash']),parser_version)).fetchone()
    if row and row['status']=='POSTED': return row['id'],'POSTED'
    stamp=now()
    category=recognition_category(p,result)
    status='BLOCKED' if result.get('errors') else 'READY'
    rows=result.get('rows',[])
    segments=sorted({clean(r.get('package')) for r in rows if clean(r.get('package'))})
    batches=sorted({clean(r.get('batch')) for r in rows if clean(r.get('batch'))})
    snapshot={k:result.get(k) for k in ('headers','sheets','sheet','sheet_kind','header_row','mapping','rows','errors','warnings','count','hash','source_rows','layout_rows','candidates','selections','reconciliation')}
    snapshot['corrections']=p.get('corrections',{})
    snapshot['input_settings']=[{k:v for k,v in s.items() if k not in ('content','text','rows')} for s in p.get('inputs',[])]
    snapshot['kind']=p.get('kind','IN')
    snapshot['defaults']=p.get('defaults',{})
    business_date=clean(p.get('arrival_date') or p.get('batch_date'))
    original_name=' + '.join(dict.fromkeys(clean(s.get('filename')) for s in p.get('inputs',[]) if s.get('filename')))
    values=(stamp,stamp,business_date,category,status,original_name or clean(p.get('original_name') or p.get('filename')) or '未命名文件',raw_hash,
            result.get('selection_hash',result['hash']),result.get('sheet',''),parser_version,source_path,
            json.dumps(segments,ensure_ascii=False),json.dumps(batches,ensure_ascii=False),
            json.dumps(snapshot,ensure_ascii=False))
    if row:
        c.execute('''UPDATE recognitions SET updated_at=?,business_date=?,category=?,status=?,original_name=?,raw_sha256=?,
                     selection_hash=?,sheet_name=?,parser_version=?,source_path=?,segments_json=?,batches_json=?,snapshot_json=? WHERE id=?''',
                  (stamp,*values[2:],row['id']))
        return row['id'],status
    cur=c.execute('''INSERT INTO recognitions(recognized_at,updated_at,business_date,category,status,original_name,raw_sha256,
                   selection_hash,sheet_name,parser_version,source_path,segments_json,batches_json,snapshot_json)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',values)
    return cur.lastrowid,status

def preview_post(p):
    result,raw,archive_name=parse_intake(p)
    p=dict(p,filename=archive_name)
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        recognition_id,status=upsert_recognition(c,p,result,raw)
        result['recognition_id']=recognition_id
        result['recognition_status']=status
    return result

def recognition_date_bound(value):
    day=date.fromisoformat(value)
    return datetime.combine(day,time.min,RECOGNITION_TIMEZONE).astimezone(timezone.utc).isoformat(timespec='seconds')

def recognition_display(record, snapshot):
    if record['original_name'].endswith('.intake.json'):
        settings=snapshot.get('input_settings') or []
        names=list(dict.fromkeys(clean(s.get('filename')) for s in settings if isinstance(s,dict) and s.get('filename')))
        mode=next((s.get('source') for s in settings if isinstance(s,dict)),None)
        record['original_name']=' + '.join(names) or {'manual':'手动录入','paste':'粘贴内容'}.get(mode,record['original_name'])
    record['can_reopen']=record['has_source'] and record['status'] in ('READY','BLOCKED') and snapshot.get('kind')=='IN'
    errors=snapshot.get('errors') or []
    record['issue_count']=len(errors)
    record['issue_summary']=clean(errors[0]) if errors else ''
    return record


def recognition_list(filters):
    start=clean(filters.get('from',[''])[0])
    end=clean(filters.get('to',[''])[0])
    status=clean(filters.get('status',[''])[0]).upper()
    category=clean(filters.get('category',[''])[0]).lower()
    segment=clean(filters.get('segment',[''])[0])
    batch=clean(filters.get('batch',[''])[0])
    try: limit=max(1,min(int(filters.get('limit',['100'])[0]),200))
    except (ValueError,TypeError): limit=100
    try: offset=max(0,int(filters.get('offset',['0'])[0]))
    except (ValueError,TypeError): offset=0
    clauses=[]; params=[]
    if start: clauses.append('r.recognized_at>=?'); params.append(recognition_date_bound(start))
    if end:
        end_exclusive=(date.fromisoformat(end)+timedelta(days=1)).isoformat()
        clauses.append('r.recognized_at<?'); params.append(recognition_date_bound(end_exclusive))
    if status: clauses.append('r.status=?'); params.append(status)
    if category: clauses.append('r.category=?'); params.append(category)
    where=(' WHERE '+' AND '.join(clauses)) if clauses else ''
    with connect() as c:
        rows=c.execute('''SELECT r.id,r.recognized_at,r.updated_at,r.business_date,r.category,r.status,r.original_name,r.raw_sha256,
                          r.sheet_name,r.source_path,r.segments_json,r.batches_json,r.snapshot_json,
                          CASE WHEN d.id IS NULL THEN NULL ELSE r.document_id END document_id
                          FROM recognitions r LEFT JOIN documents d ON d.id=r.document_id'''+where+''' ORDER BY r.recognized_at DESC,r.id DESC''',params).fetchall()
    records=[]
    for row in rows:
        r=dict(row); r['segments']=json.loads(r.pop('segments_json') or '[]'); r['batches']=json.loads(r.pop('batches_json') or '[]')
        r['has_source']=bool(r.pop('source_path'))
        snapshot=json.loads(r.pop('snapshot_json') or '{}')
        if segment and segment not in r['segments']: continue
        if batch and batch not in r['batches']: continue
        records.append(recognition_display(r,snapshot))
    return {'rows':records[offset:offset+limit],'total':len(records),'offset':offset,'limit':limit,'has_more':offset+limit<len(records)}

def recognition_detail(recognition_id):
    with connect() as c:
        row=c.execute('SELECT r.*,d.id linked_document_exists FROM recognitions r LEFT JOIN documents d ON d.id=r.document_id WHERE r.id=?',(int(recognition_id),)).fetchone()
    if not row: raise ValueError('识别归档不存在 / Recognition archive not found')
    record=dict(row); record['segments']=json.loads(record.pop('segments_json') or '[]'); record['batches']=json.loads(record.pop('batches_json') or '[]')
    if not record.pop('linked_document_exists',None): record['document_id']=None
    record['snapshot']=json.loads(record.pop('snapshot_json'))
    record['has_source']=bool(record.get('source_path'))
    return recognition_display(record,record['snapshot'])

def recognition_reopen(recognition_id):
    record=recognition_detail(recognition_id)
    if record['status'] not in ('READY','BLOCKED'):
        raise ValueError('只有未确认的归档可以继续处理 / Only unposted archives can be reopened')
    rel=record.get('source_path') or ''
    source_root=(DATA/'sources').resolve()
    path=(DATA/rel).resolve()
    try: path.relative_to(source_root)
    except ValueError: raise ValueError('原件路径无效 / Invalid source path') from None
    if not path.is_file() or path.stat().st_size>17*1024*1024:
        raise ValueError('归档原件不存在或超过处理上限 / Source is missing or exceeds the intake limit')
    try: saved=json.loads(path.read_text(encoding='utf-8'))
    except (UnicodeDecodeError,json.JSONDecodeError): raise ValueError('这条归档无法继续处理，请从入库办理重新上传原件 / Archive cannot be reopened; upload the source again') from None
    if saved.get('format')!='warehouse-intake-v1' or not isinstance(saved.get('sources'),list):
        raise ValueError('这条归档无法继续处理，请从入库办理重新上传原件 / Archive cannot be reopened; upload the source again')
    inputs=saved['sources']
    settings=record['snapshot'].get('input_settings') or []
    for i,source in enumerate(inputs):
        if not isinstance(source,dict): raise ValueError('归档来源格式无效 / Invalid archived input')
        if i<len(settings) and isinstance(settings[i],dict): source.update(settings[i])
    for selection in record['snapshot'].get('selections') or []:
        index=selection.get('input')
        if isinstance(index,int) and 0<=index<len(inputs):
            inputs[index].setdefault('sheet',selection.get('sheet'))
            inputs[index].setdefault('header',selection.get('header_row'))
    if not 1<=len(inputs)<=4: raise ValueError('归档没有可恢复的录入来源 / No restorable inputs in archive')
    original_name=record['original_name']
    if original_name.endswith('.intake.json'):
        original_name=next((s.get('filename') for s in inputs if s.get('filename')),'未命名文件')
    snapshot=record['snapshot']
    return {'id':record['id'],'inputs':inputs,'corrections':snapshot.get('corrections') or {},
            'defaults':snapshot.get('defaults') or {},'kind':snapshot.get('kind') or 'IN',
            'arrival_date':record['business_date'],'original_name':original_name}

def cancel_recognition(p):
    rec_id=int(p.get('id') or 0)
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT status FROM recognitions WHERE id=?',(rec_id,)).fetchone()
        if not row: raise ValueError('识别归档不存在 / Recognition archive not found')
        if row['status']=='POSTED': raise ValueError('已入账记录不能取消 / Posted recognition cannot be cancelled')
        c.execute("UPDATE recognitions SET status='CANCELLED',updated_at=? WHERE id=?",(now(),rec_id))
    return {'ok':True,'id':rec_id,'status':'CANCELLED'}

def delete_recognition(p):
    result=delete_recognitions(dict(p,ids=[p.get('id')]))
    return dict(result,id=result['deleted'][0])

def delete_recognitions(p):
    verify_admin_password(p.get('password'))
    ids=p.get('ids')
    if not isinstance(ids,list) or not 1<=len(ids)<=1000:
        raise ValueError('请选择1至1000条识别记录 / Select 1–1000 recognition records')
    if any(isinstance(i,bool) or not str(i).isdigit() or int(i)<=0 for i in ids):
        raise ValueError('识别记录编号无效 / Invalid recognition ID')
    ids=list(dict.fromkeys(int(i) for i in ids))
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        for rec_id in ids:
            old=c.execute('SELECT * FROM recognitions WHERE id=?',(rec_id,)).fetchone()
            if not old: raise ValueError(f'识别归档 #{rec_id} 不存在，本次未删除 / Recognition archive not found; nothing deleted')
            c.execute('INSERT INTO admin_actions(action,document_id,document_number,reason,operator,created) VALUES(?,?,?,?,?,?)',
                      ('RECOGNITION_DELETE',old['document_id'] or 0,f"识别归档 #{rec_id}：{old['original_name']}",json.dumps(dict(old),ensure_ascii=False),'密码验证',now()))
            c.execute('DELETE FROM recognitions WHERE id=?',(rec_id,))
    # Original files can be shared by other archives and posted receipts.
    return {'ok':True,'deleted':ids,'stock_changed':False}

def newdoc(c,p,kind,file_hash=None,source_path=None,reverse_of=None):
    key=required(p.get('request_key'),'请求编号 / Request key')
    row=c.execute('SELECT id FROM documents WHERE request_key=?',(key,)).fetchone()
    if row: return row['id'],True
    operator=required(p.get('operator'),'经办人 / Operator')
    no=kind+'-'+datetime.now().strftime('%Y%m%d')+'-'+secrets.token_hex(4).upper()
    # A source file and system-generated document number are sufficient for a
    # receipt.  Keep a manually entered delivery note when supplied; otherwise
    # use the unique system number so the operator is not forced to invent a
    # second reference.
    ref=required(clean(p.get('external_ref')) or no,'业务单号 / Reference')
    cur=c.execute('INSERT INTO documents(number,kind,external_ref,operator,party,purpose,created,arrival_date,request_key,file_hash,source_path,reverse_of) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(no,kind,ref,operator,clean(p.get('party')),clean(p.get('purpose')),now(),clean(p.get('arrival_date')),key,file_hash,source_path,reverse_of))
    return cur.lastrowid,False

def item(c,r):
    old=c.execute('SELECT * FROM items WHERE code=?',(r['code'],)).fetchone()
    if old:
        if any(old[k]!=r[k] for k in ('name','spec','unit')): raise ValueError('同编码名称/规格/单位冲突 / Item identity conflict')
    else: c.execute('INSERT INTO items(code,name,spec,unit,package,attribute) VALUES(?,?,?,?,?,?)',tuple(r.get(k,'') for k in ('code','name','spec','unit','package','attribute')))
def position(c,code,batch,box,warehouse,bin,state,contents=None):
    if isinstance(contents,str): contents=json.loads(contents or '[]')
    if contents is not None and not isinstance(contents,list): raise ValueError('箱内明细格式无效')
    encoded=json.dumps(contents or [],ensure_ascii=False)
    c.execute('INSERT OR IGNORE INTO stock(code,batch,box,warehouse,bin,state,qty,contents) VALUES(?,?,?,?,?,?,0,?)',(code,batch,box,warehouse,bin,state,encoded))
    row=c.execute('SELECT * FROM stock WHERE code=? AND batch=? AND box=? AND warehouse=? AND bin=? AND state=?',(code,batch,box,warehouse,bin,state)).fetchone()
    if contents and clean(row['contents']) in ('','[]'):
        c.execute('UPDATE stock SET contents=? WHERE id=?',(encoded,row['id']))
        row=c.execute('SELECT * FROM stock WHERE id=?',(row['id'],)).fetchone()
    return row
def movement(c,doc,sid,delta):
    cur=c.execute('UPDATE stock SET qty=qty+? WHERE id=? AND qty+?>=0',(delta,sid,delta))
    if cur.rowcount!=1: raise ValueError('库存不足，操作未生效 / Insufficient stock; nothing posted')
    c.execute('INSERT INTO movements(doc_id,stock_id,delta) VALUES(?,?,?)',(doc,sid,delta))
def line(c,doc,r,qty):
    c.execute('INSERT INTO lines(doc_id,code,name,spec,unit,batch,box,warehouse,bin,state,qty,attribute,contents,source_row) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(doc,*(r.get(k,'') for k in ('code','name','spec','unit','batch','box','warehouse','bin','state')),qty,r.get('attribute',''),json.dumps(r.get('contents',[]),ensure_ascii=False) if not isinstance(r.get('contents',''),str) else r.get('contents',''),r.get('source_row')))

def current_design_qty(c,code):
    row=c.execute('SELECT qty FROM design WHERE code=?',(code,)).fetchone()
    if not row: return None
    change=c.execute('SELECT COALESCE(SUM(delta),0) FROM changes WHERE code=?',(code,)).fetchone()[0]
    return int(row['qty'])+int(change or 0)

def validate_receipt_limits(c,rows):
    """Prevent a confirmed receipt from exceeding a known design quantity.

    A missing design baseline remains allowed for the current arrival-only
    workflow; once a baseline exists, non-spare receipts cannot exceed it.
    This check happens before the document is created so the transaction
    remains all-or-nothing.
    """
    incoming={}
    for row in rows:
        if row.get('state')=='SPARE': continue
        incoming[row['code']]=incoming.get(row['code'],0)+int(row.get('amount',0))
    for code,amount in incoming.items():
        limit=current_design_qty(c,code)
        if limit is None: continue
        received=c.execute("SELECT COALESCE(SUM(l.qty),0) FROM lines l JOIN documents d ON l.doc_id=d.id WHERE l.code=? AND d.kind='IN' AND d.reversed_by IS NULL AND l.state!='SPARE'",(code,)).fetchone()[0]
        if int(received or 0)+amount>limit:
            raise ValueError(f'{code}: 入库数量将超过合同/设计数量（已有 {number(received)}，本次 {number(amount)}，上限 {number(limit)}） / Receipt would exceed contract/design quantity (received {number(received)}, this receipt {number(amount)}, limit {number(limit)})')

def import_post(p):
    result,raw,archive_name=parse_intake(p)
    p=dict(p,filename=archive_name)
    kind=p.get('kind','IN')
    if kind not in ('IN','BASELINE'): raise ValueError('无效导入类型 / Invalid import type')
    sheet_label=result['sheet'].lower()
    # Reject a carton-level shipment worksheet before item-identity checks.
    # This keeps the operator-facing message about choosing the packing-detail
    # sheet even when the same workbook was already used for a receipt.
    if result['sheet']=='CSV': sheet_label=f"{sheet_label} {p.get('filename','')}".lower()
    if kind=='IN' and 'inputs' not in p and ('发货' in sheet_label or 'shipment' in sheet_label or 'carton' in sheet_label):
        raise ValueError('发货清单是箱级物流核对表，请选择装箱清单办理入库 / Shipment list is for carton-level logistics reconciliation; select the packing-detail sheet for receipt posting')
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        prior=c.execute('SELECT id FROM documents WHERE request_key=?',(p.get('request_key'),)).fetchone()
        if prior: return {'id':prior['id']}
        recognition_id,_=upsert_recognition(c,p,result,raw)
        submitted_id=p.get('recognition_id')
        if submitted_id is not None and int(submitted_id)!=recognition_id:
            raise ValueError('识别归档与当前文件不匹配，请重新预览 / Recognition archive does not match this file; preview again')
        if result['errors']: raise ValueError('\n'.join(result['errors'][:12]))
        if kind=='IN': validate_receipt_limits(c,result['rows'])
        _,path=store_source(raw,p['filename'])
        doc,_=newdoc(c,{**p,'arrival_date':p.get('arrival_date') or p.get('batch_date') or ''},kind,result['hash'],path)
        for r in result['rows']:
            item(c,r)
            if kind=='BASELINE':
                if c.execute('SELECT 1 FROM design WHERE code=?',(r['code'],)).fetchone(): raise ValueError('设计量已锁定，请使用设计变更 / Baseline locked; use change record')
                c.execute('INSERT INTO design VALUES(?,?)',(r['code'],r['amount']))
            else:
                s=position(c,r['code'],r['batch'],r['box'],r['warehouse'],r['bin'],r['state'],r.get('contents'))
                movement(c,doc,s['id'],r['amount'])
                c.execute('INSERT OR IGNORE INTO batches VALUES(?,?,?,?)',(r['batch'],clean(p.get('party')),'ARRIVED',now()))
            line(c,doc,r,r['amount'])
        c.execute("UPDATE recognitions SET status='POSTED',document_id=?,updated_at=? WHERE id=?",(doc,now(),recognition_id))
    return {'id':doc}

def stock_post(p):
    kind=p.get('kind','OUT')
    if kind not in ('OUT','HANDOVER','MOVE','STATE'): raise ValueError('业务类型无效 / Invalid transaction')
    required(p.get('party'),'领用/接收单位 / Receiving party')
    required(p.get('purpose'),'用途/原因 / Purpose')
    rows=p.get('lines',[])
    if not rows or len(rows)>100: raise ValueError('选择1至100条物资 / Select 1–100 lines')
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        doc,exists=newdoc(c,p,kind)
        if exists: return {'id':doc}
        for r in rows:
            s=c.execute('SELECT s.*,i.name,i.spec,i.unit FROM stock s JOIN items i ON i.code=s.code WHERE s.id=?',(int(r['stock_id']),)).fetchone()
            if not s: raise ValueError('库位记录不存在 / Position not found')
            s=dict(s); q=units(r['qty'])
            remark=r.get('remark','')
            if remark is None: remark=''
            if not isinstance(remark,str) or len(remark)>1000: raise ValueError('备注须为文字且不超过1000字')
            # Outgoing line attributes are transaction notes, separate from item attributes.
            s['attribute']=clean(remark)
            if kind=='OUT' and s['state']!='AVAILABLE': raise ValueError('仅可用库存可领用 / Only available stock may be issued')
            if kind=='HANDOVER' and s['state']!='SPARE': raise ValueError('请先预留备品 / Reserve spares before handover')
            movement(c,doc,s['id'],-q); line(c,doc,s,-q)
            if kind in ('MOVE','STATE'):
                state=p.get('state') if kind=='STATE' else s['state']
                if state not in STATES: raise ValueError('状态无效 / Invalid state')
                wh=required(p.get('warehouse'),'仓库 / Warehouse') if kind=='MOVE' else s['warehouse']
                loc=required(p.get('bin'),'库位 / Bin') if kind=='MOVE' else s['bin']
                target=position(c,s['code'],s['batch'],s['box'],wh,loc,state,s.get('contents'))
                if target['id']==s['id']: raise ValueError('目标与原记录相同 / Same destination')
                movement(c,doc,target['id'],q)
                line(c,doc,dict(s,warehouse=wh,bin=loc,state=state),q)
    return {'id':doc}

def reverse(p):
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        old=c.execute('SELECT * FROM documents WHERE id=?',(int(p['id']),)).fetchone()
        if not old or old['kind'] in ('BASELINE','REV') or old['reversed_by']: raise ValueError('此单不可冲销 / Cannot reverse this document')
        required(p.get('purpose'),'冲销原因 / Reversal reason')
        doc,exists=newdoc(c,p,'REV',reverse_of=old['id'])
        if exists: return {'id':doc}
        for m in c.execute('SELECT * FROM movements WHERE doc_id=? ORDER BY delta ASC',(old['id'],)).fetchall(): movement(c,doc,m['stock_id'],-m['delta'])
        for r in c.execute('SELECT * FROM lines WHERE doc_id=?',(old['id'],)).fetchall(): line(c,doc,dict(r),-r['qty'])
        c.execute('UPDATE documents SET reversed_by=? WHERE id=?',(doc,old['id']))
    return {'id':doc}

def admin_password_path():
    return DATA/'admin_password.sha256'

def admin_password_configured():
    if clean(os.environ.get('MATERIALS_ADMIN_PASSWORD')):
        return True
    try:
        return bool(clean(admin_password_path().read_text(encoding='utf-8')))
    except (FileNotFoundError, OSError):
        return False

def verify_admin_password(value):
    supplied=clean(value)
    if not supplied: raise ValueError('请输入管理员密码 / Enter the administrator password')
    configured=clean(os.environ.get('MATERIALS_ADMIN_PASSWORD'))
    if configured:
        ok=hmac.compare_digest(supplied,configured)
    else:
        try: expected=clean(admin_password_path().read_text(encoding='utf-8'))
        except (FileNotFoundError, OSError): expected=''
        ok=bool(expected) and hmac.compare_digest(hashlib.sha256(supplied.encode('utf-8')).hexdigest(),expected)
    if not ok: raise ValueError('管理员密码错误 / Incorrect administrator password')

def admin_password_set(p):
    new=required(p.get('new_password'),'新管理员密码 / New administrator password',64)
    if clean(os.environ.get('MATERIALS_ADMIN_PASSWORD')):
        raise ValueError('密码由本机启动配置管理，请修改启动配置 / Password is managed by local startup configuration')
    if admin_password_configured(): verify_admin_password(p.get('current_password'))
    DATA.mkdir(parents=True,exist_ok=True)
    admin_password_path().write_text(hashlib.sha256(new.encode('utf-8')).hexdigest(),encoding='utf-8')
    return {'ok':True,'configured':True}

def cleanup_orphan_item(c,code):
    # Keep the item master when it is still referenced by another receipt,
    # design baseline/change, or a non-zero stock position.  A mistaken-only
    # import can therefore disappear from the visible materials table after
    # its document is removed without risking a real balance.
    if c.execute('SELECT 1 FROM lines WHERE code=? LIMIT 1',(code,)).fetchone(): return
    if c.execute('SELECT 1 FROM design WHERE code=? LIMIT 1',(code,)).fetchone(): return
    if c.execute('SELECT 1 FROM changes WHERE code=? LIMIT 1',(code,)).fetchone(): return
    if c.execute('SELECT 1 FROM stock WHERE code=? AND qty<>0 LIMIT 1',(code,)).fetchone(): return
    c.execute('DELETE FROM stock WHERE code=? AND NOT EXISTS(SELECT 1 FROM movements WHERE stock_id=stock.id)',(code,))
    c.execute('DELETE FROM items WHERE code=? AND NOT EXISTS(SELECT 1 FROM lines WHERE code=?) AND NOT EXISTS(SELECT 1 FROM stock WHERE code=?)',(code,code,code))

def admin_delete_in_transaction(c,doc_id,reason,operator):
    old=c.execute('SELECT * FROM documents WHERE id=?',(int(doc_id),)).fetchone()
    if not old: raise ValueError('单据不存在 / Document not found')
    if old['kind'] in ('BASELINE','REV') or old['reversed_by']:
        raise ValueError('设计基准单、冲销单或已冲销单不能删除 / Baseline, reversal or already reversed documents cannot be deleted')
    affected_codes={row['code'] for row in c.execute('SELECT code FROM lines WHERE doc_id=?',(old['id'],)).fetchall()}
    movements=c.execute('SELECT * FROM movements WHERE doc_id=? ORDER BY delta DESC,id DESC',(old['id'],)).fetchall()
    # Undo each movement in reverse posting direction.  If a receipt has
    # already been consumed by a later issue, refuse the hard delete and
    # direct the operator to the auditable reversal action instead.
    for m in movements:
        if c.execute('SELECT 1 FROM movements WHERE stock_id=? AND id>? AND doc_id<>? LIMIT 1',(m['stock_id'],m['id'],old['id'])).fetchone():
            raise ValueError('该单据已有后续库存流水，请使用作废/冲销保留流水 / Later stock movements exist; use void/reversal to keep the audit trail')
        cur=c.execute('UPDATE stock SET qty=qty-? WHERE id=? AND qty-?>=0',(m['delta'],m['stock_id'],m['delta']))
        if cur.rowcount!=1:
            raise ValueError('该单据已影响后续库存，请使用作废/冲销保留流水 / Later transactions depend on this document; use void/reversal to keep the audit trail')
    c.execute('INSERT INTO admin_actions(action,document_id,document_number,reason,operator,created) VALUES(?,?,?,?,?,?)',('DELETE',old['id'],old['number'],reason,operator,now()))
    c.execute('DELETE FROM movements WHERE doc_id=?',(old['id'],))
    c.execute('DELETE FROM lines WHERE doc_id=?',(old['id'],))
    c.execute('DELETE FROM documents WHERE id=?',(old['id'],))
    c.execute("UPDATE recognitions SET status='CANCELLED',document_id=NULL,updated_at=? WHERE document_id=?",(now(),old['id']))
    for code in affected_codes: cleanup_orphan_item(c,code)
    return old['number']

def admin_delete(p):
    verify_admin_password(p.get('password'))
    reason=clean(p.get('purpose')) or '密码确认删除'
    operator=clean(p.get('operator')) or '密码验证'
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        deleted=admin_delete_in_transaction(c,p['id'],reason,operator)
    return {'ok':True,'deleted':deleted}

def admin_delete_batch(p):
    verify_admin_password(p.get('password'))
    reason=clean(p.get('purpose')) or '密码确认删除'
    operator=clean(p.get('operator')) or '密码验证'
    ids=p.get('ids') or []
    if not isinstance(ids,list) or not ids or len(ids)>100: raise ValueError('请选择1至100张单据 / Select 1–100 documents')
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        deleted=[admin_delete_in_transaction(c,doc_id,reason,operator) for doc_id in ids]
    return {'ok':True,'deleted':deleted}

def stock_balance_fields(c,stocks):
    """Attach one server-calculated balance summary to every visible stock row."""
    receipt_totals={}
    for row in c.execute("""
        SELECT l.code,l.batch,l.box,i.package,COALESCE(SUM(l.qty),0) total
        FROM lines l
        JOIN documents d ON d.id=l.doc_id
        JOIN items i ON i.code=l.code
        WHERE d.kind='IN' AND d.reversed_by IS NULL
        GROUP BY l.code,l.batch,l.box,i.package
    """):
        receipt_totals[(row['code'],row['batch'],row['box'],row['package'])]=int(row['total'] or 0)
    by_state={}
    for row in c.execute("""
        SELECT s.code,s.batch,s.box,i.package,s.state,COALESCE(SUM(s.qty),0) total
        FROM stock s JOIN items i ON i.code=s.code
        GROUP BY s.code,s.batch,s.box,i.package,s.state
    """):
        key=(row['code'],row['batch'],row['box'],row['package'])
        by_state.setdefault(key,{})[row['state']]=int(row['total'] or 0)
    for row in stocks:
        key=(row['code'],row['batch'],row['box'],row['package'])
        state_totals=by_state.get(key,{})
        row['received_total']=number(receipt_totals[key]) if key in receipt_totals else None
        row['available_qty']=number(state_totals.get('AVAILABLE',0))
        row['pending_qty']=number(state_totals.get('PENDING',0))
        row['quarantine_qty']=number(state_totals.get('QUARANTINE',0))
        row['spare_qty']=number(state_totals.get('SPARE',0))
    return stocks

def reconciliation(c,items):
    """Compare movement totals with stock positions and known design limits."""
    movement_totals={}
    for row in c.execute("""
        SELECT s.code,i.package,COALESCE(SUM(m.delta),0) total
        FROM movements m JOIN stock s ON s.id=m.stock_id JOIN items i ON i.code=s.code
        GROUP BY s.code,i.package
    """):
        movement_totals[(row['code'],row['package'])]=int(row['total'] or 0)
    stock_totals={}
    for row in c.execute("""
        SELECT s.code,i.package,COALESCE(SUM(s.qty),0) total
        FROM stock s JOIN items i ON i.code=s.code
        GROUP BY s.code,i.package
    """):
        stock_totals[(row['code'],row['package'])]=int(row['total'] or 0)
    issues=[]
    # 先建索引，避免每条异常都遍历整个物资列表。
    by_code={r.get('code'):r for r in items}
    for key in sorted(set(movement_totals)|set(stock_totals)):
        movement_total=movement_totals.get(key,0); stock_total=stock_totals.get(key,0)
        if movement_total!=stock_total:
            code,package=key
            source=by_code.get(code,{})
            name=source.get('name','')
            unit=source.get('unit','')
            issues.append({'kind':'ledger_mismatch','code':code,'name':name,'package':package,'unit':unit,'expected':number(movement_total),'actual':number(stock_total),'difference':number(stock_total-movement_total)})
    for row in items:
        if row.get('current') is None or row.get('received') is None: continue
        if float(row['received'])>float(row['current'])+1e-9:
            issues.append({'kind':'over_contract','code':row['code'],'name':row['name'],'package':row.get('package',''),'unit':row.get('unit',''),'expected':row['current'],'actual':row['received'],'difference':row['received']-row['current']})
    return {'ok':not issues,'issue_count':len(issues),'issues':issues}

LEDGER_SQL = '''
    SELECT m.id AS movement_id, m.delta, d.id AS doc_id, d.number, d.kind,
           d.external_ref, d.created, d.arrival_date, d.operator, d.party,
           d.purpose, (d.source_path<>'') AS has_source, s.code, i.name, i.spec, i.unit, i.package,
           s.batch, s.box, s.warehouse, s.bin, s.state
    FROM movements m
    JOIN documents d ON d.id=m.doc_id
    JOIN stock s ON s.id=m.stock_id
    JOIN items i ON i.code=s.code
'''
LEDGER_KINDS = ('IN', 'OUT', 'HANDOVER', 'MOVE', 'STATE', 'REV', 'BASELINE')

def ledger_filter(kind, q):
    """Build the WHERE clause shared by the ledger count and page queries."""
    where, args = [], []
    if kind:
        if kind not in LEDGER_KINDS: raise ValueError('流水类型无效 / Invalid ledger kind')
        where.append('d.kind=?'); args.append(kind)
    if q:
        # Chinese text has no case, LOWER() is only meaningful for codes/refs.
        like = '%' + q.lower()[:100] + '%'
        columns = ('d.number', 'd.external_ref', 'd.party', 'd.purpose', 'd.operator',
                   's.code', 's.batch', 's.box', 's.warehouse', 's.bin',
                   'i.name', 'i.spec', 'i.package', 'i.unit', 's.state')
        where.append('(' + ' OR '.join(f'LOWER({col}) LIKE ?' for col in columns) + ')')
        args.extend([like] * len(columns))
    return (' WHERE ' + ' AND '.join(where) if where else ''), args

def ledger_page(c, limit=LEDGER_PAGE_SIZE, offset=0, kind='', q='', stock_id=None):
    """Return one page of movements, newest first, plus the filtered total."""
    clause, args = ledger_filter(kind, q)
    if stock_id is not None:
        clause += (' AND ' if clause else ' WHERE ') + 'EXISTS (SELECT 1 FROM stock target WHERE target.id=? AND s.code=target.code AND s.batch=target.batch AND s.box=target.box)'
        args.append(stock_id)
    limit = max(1, min(int(limit), LEDGER_MAX_PAGE_SIZE))
    offset = max(0, int(offset))
    total = c.execute('SELECT COUNT(*) FROM movements m '
                      'JOIN documents d ON d.id=m.doc_id '
                      'JOIN stock s ON s.id=m.stock_id '
                      'JOIN items i ON i.code=s.code' + clause, args).fetchone()[0]
    rows = []
    for r in c.execute(LEDGER_SQL + clause + ' ORDER BY m.id DESC LIMIT ? OFFSET ?', args + [limit, offset]):
        row = dict(r); row['delta'] = number(row['delta']); row['qty'] = abs(row['delta'])
        row['direction'] = 'IN' if row['delta'] > 0 else 'OUT'
        row['date'] = row['arrival_date'] or row['created']
        rows.append(row)
    return {'rows': rows, 'total': total, 'offset': offset, 'limit': limit,
            'has_more': offset + len(rows) < total}

def build_items(c):
    """Every item with its computed totals.

    Shared by the overview page, the CSV export and the paged item table so all
    three always report identical numbers.  The aggregates are fetched in four
    batch queries; the previous per-item version ran 7 SQL statements per item
    (35,000 statements at 5,000 items) and took over five minutes per refresh.
    """
    change_map={row['code']:row['total'] for row in c.execute(
        'SELECT code,COALESCE(SUM(delta),0) total FROM changes GROUP BY code')}
    # 到货量要排除备品，出库与移交不排除，所以按「是否备品」分组一次取全。
    line_agg={}
    for row in c.execute("SELECT l.code,d.kind,CASE WHEN l.state='SPARE' THEN 1 ELSE 0 END spare,"
                         "COALESCE(SUM(l.qty),0) total FROM lines l JOIN documents d ON l.doc_id=d.id "
                         "WHERE d.reversed_by IS NULL GROUP BY l.code,d.kind,spare"):
        line_agg[(row['code'],row['kind'],row['spare'])]=row['total']
    date_map={}
    for row in c.execute("SELECT l.code,d.kind,"
                         "MIN(CASE WHEN d.kind='IN' AND d.arrival_date<>'' THEN d.arrival_date ELSE d.created END) first_date,"
                         "MAX(CASE WHEN d.kind='IN' AND d.arrival_date<>'' THEN d.arrival_date ELSE d.created END) last_date "
                         "FROM lines l JOIN documents d ON l.doc_id=d.id "
                         "WHERE d.kind IN ('IN','OUT') AND d.reversed_by IS NULL GROUP BY l.code,d.kind"):
        date_map[(row['code'],row['kind'])]=(row['first_date'],row['last_date'])
    stock_map={row['code']:(row['onhand'],row['spare']) for row in c.execute(
        "SELECT code,COALESCE(SUM(CASE WHEN state<>'SPARE' THEN qty ELSE 0 END),0) onhand,"
        "COALESCE(SUM(CASE WHEN state='SPARE' THEN qty ELSE 0 END),0) spare FROM stock GROUP BY code")}
    items=[]
    for r in c.execute('SELECT i.*,d.qty baseline FROM items i LEFT JOIN design d ON d.code=i.code ORDER BY i.code'):
        r=dict(r); code=r['code']
        change=change_map.get(code,0)
        r['change']=number(change); r['current']=number(r['baseline']+change) if r['baseline'] is not None else None
        r['baseline']=number(r['baseline']) if r['baseline'] is not None else None
        for name,kind in [('received','IN'),('issued','OUT'),('handed','HANDOVER')]:
            # Reserved spares are tracked separately from the contract's main
            # design quantity and must not inflate the arrived/remaining view.
            non_spare=line_agg.get((code,kind,0),0)
            q=non_spare if name=='received' else non_spare+line_agg.get((code,kind,1),0)
            r[name]=number(abs(q))
        for kind,prefix in [('IN','inbound'),('OUT','outbound')]:
            first_date,last_date=date_map.get((code,kind),('',''))
            r[prefix+'_first']=first_date or ''; r[prefix+'_last']=last_date or ''
        onhand,spare=stock_map.get(code,(0,0))
        r['onhand']=number(onhand); r['spare']=number(spare)
        r['outstanding']=max(0,r['current']-r['received']) if r['current'] is not None else None
        items.append(r)
    return items


def blank_contract(code):
    return {'code':code,'item_codes':[],'item_count':0,'baseline':None,'change':0,'current':None,
            'received':0,'issued':0,'handed':0,'onhand':0,'spare':0,'outstanding':None,
            'unit_totals':{},'inbound_first':'','inbound_last':'','outbound_first':'','outbound_last':''}


def build_contracts(items):
    """Roll items up into the fixed segments.

    Carries item codes plus a count instead of a second copy of every item row.
    That nested copy used to double the response size; the segment page now
    resolves the codes against the paged item table instead.
    """
    contracts={}
    for r in items:
        key=clean(r.get('package')) or '未指定标段'
        g=contracts.setdefault(key,blank_contract(key))
        g['item_codes'].append(r['code']); g['item_count']+=1
        for k in ('baseline','change','current','received','issued','handed','onhand','spare','outstanding'):
            if r[k] is not None:
                g[k]=r[k] if g[k] is None else g[k]+r[k]
        unit=r.get('unit') or '—'; totals=g['unit_totals'].setdefault(unit,{'current':None,'received':0,'issued':0,'onhand':0,'spare':0,'outstanding':0})
        for k in ('current','received','issued','onhand','spare','outstanding'):
            if r[k] is not None: totals[k]=r[k] if totals[k] is None else totals[k]+r[k]
        for k in ('inbound_first','inbound_last','outbound_first','outbound_last'):
            val=r.get(k) or ''
            if val and (not g[k] or (k.endswith('_first') and val<g[k]) or (k.endswith('_last') and val>g[k])): g[k]=val
    # Keep the fixed project segments visible on the leadership view even
    # before their first receipt is posted.
    for code, _, _ in SEGMENTS:
        contracts.setdefault(code,blank_contract(code))
    return list(contracts.values())


def item_counts(items):
    """Headline counters, so the materials page does not need every row loaded."""
    return {'total':len(items),
            'in_stock':sum(1 for r in items if number(r.get('onhand') or 0)>0),
            'outstanding':sum(1 for r in items if number(r.get('outstanding') or 0)>0)}


def slice_items(items,limit,offset=0):
    """One bounded page out of an already-built item list."""
    limit=max(1,min(int(limit),ITEMS_MAX_PAGE_SIZE)); offset=max(0,int(offset))
    page=items[offset:offset+limit]
    return {'rows':page,'total':len(items),'offset':offset,'limit':limit,
            'has_more':offset+len(page)<len(items)}


def items_page(c,limit=ITEMS_PAGE_SIZE,offset=0,q='',package=''):
    """Filter the whole catalogue on the server, then return a single page.

    Filtering here rather than in the browser is what makes search honest: the
    page only ever holds a slice, so a client-side filter would silently miss
    every match that has not been paged in yet.
    """
    rows=build_items(c)
    if package: rows=[r for r in rows if (clean(r.get('package')) or '未指定标段')==package]
    if q:
        needle=q.lower()[:100]
        rows=[r for r in rows if any(needle in str(r.get(col) or '').lower() for col in ITEM_SEARCH_FIELDS)]
    return slice_items(rows,limit,offset)


def project_arrivals(c):
    """Small read-only major-equipment summary; unknown matches stay null."""
    packages=tuple(sorted({item['package'] for item in PROJECT_ARRIVAL_TARGETS}))
    marks=','.join('?' for _ in packages)
    rows=[dict(row) for row in c.execute(
        f'''SELECT i.code,i.name,i.spec,i.unit,i.package,d.qty baseline,
                   COALESCE(SUM(ch.delta),0) change
            FROM items i LEFT JOIN design d ON d.code=i.code
            LEFT JOIN changes ch ON ch.code=i.code
            WHERE i.package IN ({marks}) GROUP BY i.code''', packages)]
    norm=lambda value: re.sub(r'\s+','',unicodedata.normalize('NFKC',clean(value))).casefold()
    matched_codes={}
    candidates_by_target={}
    for target in PROJECT_ARRIVAL_TARGETS:
        aliases={norm(name) for name in target['aliases']}
        matches=[row for row in rows if row['package']==target['package'] and norm(row['name']) in aliases]
        candidates_by_target[target['key']]=matches
        for row in matches: matched_codes[row['code']]=None
    receipts={}
    if matched_codes:
        codes=tuple(matched_codes)
        code_marks=','.join('?' for _ in codes)
        receipts={row['code']:row['qty'] for row in c.execute(
            f'''SELECT l.code,COALESCE(SUM(l.qty),0) qty
                FROM lines l JOIN documents d ON d.id=l.doc_id
                WHERE l.code IN ({code_marks}) AND d.kind='IN'
                  AND d.reversed_by IS NULL AND l.state<>'SPARE'
                GROUP BY l.code''', codes)}

    items=[]
    transit_reason='当前没有可追溯的数量化发运/清关记录；不代表SARS EDI实时数据'
    for target in PROJECT_ARRIVAL_TARGETS:
        lot=LOT_CROSSWALK.get(target['package'])
        matches=candidates_by_target[target['key']]
        valid=[row for row in matches if clean(row['unit'])==target['unit']]
        if not matches:
            valid=[None]
        identities={}
        for row in valid:
            if row is not None:
                identity=(norm(row['name']),norm(row['spec']),clean(row['unit']))
                identities[identity]=identities.get(identity,0)+1
        for row in valid:
            item={'view':target['view'],'equipment_key':target['key'],'name':target['name'],
                  'lot_code':lot,'spec':None,'unit':target['unit'],'received':None,
                  'in_transit':None,'total':None,'total_basis':'unknown','rate':None,
                  'status':'unmapped','reason':'未找到与设备类别、标段及单位明确匹配的物资档案',
                  'in_transit_reason':transit_reason}
            if target.get('model_display_count') is not None:
                item['model_display_count']=target['model_display_count']
            if row is None:
                items.append(item)
                continue
            item['spec']=clean(row['spec']) or None
            item['unit']=clean(row['unit']) or None
            identity=(norm(row['name']),norm(row['spec']),clean(row['unit']))
            if clean(row['unit'])!=target['unit']:
                item['status']='unit_mismatch'; item['reason']='物资档案单位与设备统计单位不一致，未纳入数量'
            elif identities.get(identity,0)>1:
                item['status']='multiple_matches'; item['reason']='相同名称、规格和单位对应多个物资编码，未合并计数'
            else:
                item['status']='mapped'; item['reason']=''
                item['received']=number(receipts.get(row['code'],0))
                if row['baseline'] is not None:
                    item['total']=number(row['baseline']+row['change'])
                    item['total_basis']='warehouse_design'
                    if item['total']>0: item['rate']=item['received']/item['total']
                    else: item['reason']='仓储设计总量为零或无效，到货率未计算'
                elif not item['reason']:
                    item['reason']='仓储设计总量待补，到货率未计算'
                if not item['spec']:
                    item['reason']=(item['reason']+'；' if item['reason'] else '')+'规格型号待设备清单确认'
            items.append(item)
        for row in matches:
            if clean(row['unit'])==target['unit']: continue
            item={'view':target['view'],'equipment_key':target['key'],'name':target['name'],
                  'lot_code':lot,'spec':clean(row['spec']) or None,'unit':clean(row['unit']) or None,
                  'received':None,'in_transit':None,'total':None,'total_basis':'unknown','rate':None,
                  'status':'unit_mismatch','reason':'物资档案单位与设备统计单位不一致，未纳入数量',
                  'in_transit_reason':transit_reason}
            if target.get('model_display_count') is not None:
                item['model_display_count']=target['model_display_count']
            items.append(item)
    return {'generated_at':now(),'items':items}


def overview():
    with connect() as c:
        stocks=[dict(r) for r in c.execute('SELECT s.*,i.name,i.spec,i.unit,i.package,i.attribute FROM stock s JOIN items i ON i.code=s.code WHERE s.qty>0 ORDER BY s.code,s.warehouse,s.bin')]
        for r in stocks:
            r['qty']=number(r['qty'])
            try: r['contents']=json.loads(r.get('contents') or '[]')
            except (TypeError,ValueError): r['contents']=[]
        stock_balance_fields(c,stocks)
        docs=[dict(r) for r in c.execute('SELECT * FROM documents ORDER BY id DESC')]
        for d in docs: d.pop('file_hash'); d.pop('request_key'); d.pop('source_path')
        # Item totals and contract roll-ups come from shared helpers: the
        # overview page, the CSV export and the paged item table must not drift
        # apart on numbers.
        items=build_items(c)
        # Only the newest slice of each unbounded list travels with the page;
        # older movements come from /api/ledger and older items from /api/items.
        page=ledger_page(c,limit=LEDGER_PAGE_SIZE)
        item_page=slice_items(items,ITEMS_PAGE_SIZE)
        pending=c.execute("SELECT COUNT(*) FROM recognitions WHERE status IN ('READY','BLOCKED') AND document_id IS NULL").fetchone()[0]
        return {'recognition_pending':pending,'items':item_page['rows'],'items_total':item_page['total'],'items_limit':item_page['limit'],'items_has_more':item_page['has_more'],'item_counts':item_counts(items),'stock':stocks,'documents':docs,'contracts':build_contracts(items),'ledger':page['rows'],'ledger_total':page['total'],'ledger_limit':page['limit'],'ledger_has_more':page['has_more'],'reconciliation':reconciliation(c,items),'batches':[dict(r) for r in c.execute('SELECT * FROM batches ORDER BY updated DESC')],'changes':[dict(r) for r in c.execute('SELECT code,delta,reason,operator,created FROM changes ORDER BY id DESC')],'events':[dict(r) for r in c.execute('SELECT * FROM events ORDER BY id DESC LIMIT 100')],'segments':[{'code':code,'zh':zh,'en':en} for code,zh,en in SEGMENTS]}

def stock_detail(stock_id, limit=100, offset=0):
    if isinstance(stock_id,bool): raise ValueError('库存位置无效')
    try: stock_id=int(stock_id)
    except (TypeError,ValueError): raise ValueError('库存位置无效')
    if stock_id<=0: raise ValueError('库存位置无效')
    with connect() as c:
        row=c.execute('SELECT s.*,i.name,i.spec,i.unit,i.package,i.attribute FROM stock s JOIN items i ON i.code=s.code WHERE s.id=?',(stock_id,)).fetchone()
        if not row: raise ValueError('库存位置不存在')
        stock=dict(row); stock['qty']=number(stock['qty'])
        try: stock['contents']=json.loads(stock.get('contents') or '[]')
        except (TypeError,ValueError): stock['contents']=[]
        stock_balance_fields(c,[stock])
        return {'stock':stock,'history':ledger_page(c,limit,offset,stock_id=stock_id)}

def doc_detail(doc):
    with connect() as c:
        d=c.execute('SELECT * FROM documents WHERE id=?',(int(doc),)).fetchone()
        if not d: raise ValueError('单据不存在 / Document not found')
        d=dict(d); d.pop('request_key'); d.pop('file_hash')
        d['lines']=[dict(r) for r in c.execute('SELECT * FROM lines WHERE doc_id=?',(doc,))]
        for r in d['lines']:
            try: r['contents']=json.loads(r.get('contents') or '[]')
            except (TypeError,ValueError): r['contents']=[]
        for r in d['lines']:
            r['qty']=number(r['qty'])
            r['remark']=r.get('attribute','') if d['kind'] in ('OUT','HANDOVER','MOVE','STATE') else ''
        return d

def change_design(p):
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        if c.execute('SELECT 1 FROM changes WHERE request_key=?',(p.get('request_key'),)).fetchone(): return {'ok':True}
        code=required(p.get('code'),'物资编码 / Code'); q=units(p.get('qty'),False)
        if q==0: raise ValueError('变更不能为零 / Change cannot be zero')
        row=c.execute('SELECT qty FROM design WHERE code=?',(code,)).fetchone()
        if not row: raise ValueError('请先导入设计量 / Import baseline first')
        total=row[0]+c.execute('SELECT COALESCE(SUM(delta),0) FROM changes WHERE code=?',(code,)).fetchone()[0]+q
        if total<0: raise ValueError('现行需求不能为负 / Requirement cannot be negative')
        c.execute('INSERT INTO changes(code,delta,reason,operator,created,request_key) VALUES(?,?,?,?,?,?)',(code,q,required(p.get('reason'),'变更依据 / Change basis'),required(p.get('operator'),'经办人 / Operator'),now(),required(p.get('request_key'),'请求编号')))
    return {'ok':True}
def batch_update(p):
    with connect() as c:
        c.execute('BEGIN IMMEDIATE')
        code=required(p.get('code'),'批次 / Batch'); status=p.get('status')
        if status not in TRANSPORT: raise ValueError('运输状态无效 / Invalid transport status')
        old=c.execute('SELECT * FROM batches WHERE code=?',(code,)).fetchone()
        c.execute('INSERT INTO batches VALUES(?,?,?,?) ON CONFLICT(code) DO UPDATE SET status=excluded.status,supplier=excluded.supplier,updated=excluded.updated',(code,clean(p.get('supplier')),status,now()))
        c.execute('INSERT INTO events(batch,old_status,new_status,operator,note,created) VALUES(?,?,?,?,?,?)',(code,old['status'] if old else None,status,required(p.get('operator'),'经办人 / Operator'),clean(p.get('note')),now()))
    return {'ok':True}

def source_download(relative, original_name='source'):
    path=(DATA/relative).resolve()
    try: path.relative_to((DATA/'sources').resolve())
    except ValueError: raise ValueError('原件路径无效 / Invalid source path') from None
    if not path.is_file() or path.stat().st_size>17*1024*1024:
        raise ValueError('原件不存在或超过处理上限 / Source missing or too large')
    raw=path.read_bytes()
    def filename(value):
        return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip(' .')[:180] or 'source'
    def mime(name):
        return {'.xls':'application/vnd.ms-excel','.xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                '.csv':'text/csv; charset=utf-8'}.get(Path(name).suffix.lower(),'application/octet-stream')
    if path.suffix.lower()!='.json':
        name=filename(original_name)
        if Path(name).suffix.lower()!=path.suffix.lower(): name+=path.suffix
        return raw,mime(name),name
    saved=json.loads(raw)
    if not isinstance(saved,dict) or saved.get('format')!='warehouse-intake-v1' or not isinstance(saved.get('sources'),list) or not 1<=len(saved['sources'])<=4:
        raise ValueError('来源记录格式无效 / Invalid source archive')
    files=[]
    import intake
    for source in saved['sources']:
        if source.get('source')=='upload':
            try: content=base64.b64decode(source.get('content',''),validate=True)
            except (ValueError,TypeError): raise ValueError('原文件编码无效 / Invalid source encoding') from None
            if not content or len(content)>12*1024*1024: raise ValueError('原文件大小无效 / Invalid source size')
            name=filename(source.get('filename') or 'source.bin')
        else:
            wb=openpyxl.Workbook(); wb.remove(wb.active)
            for title,sheet in intake.read_input(source,sys.modules[__name__]).items():
                ws=wb.create_sheet(title); ws.freeze_panes='A2'
                for row in sheet['rows']:
                    ws.append(row)
                    for cell in ws[ws.max_row]:
                        if isinstance(cell.value,str): cell.data_type='s'
                for column in ws.columns:
                    ws.column_dimensions[column[0].column_letter].width=min(45,max(14,max(len(str(c.value or '')) for c in column)*2))
            output=io.BytesIO(); wb.save(output); wb.close(); content=output.getvalue()
            name=('到货单' if source.get('role')=='arrival_note' else '装箱单')+'.xlsx'
        if (name,content) not in files: files.append((name,content))
    if len(files)==1:
        name,content=files[0]; return content,mime(name),name
    output=io.BytesIO(); used=set()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,content in files:
            original=name; suffix=2
            while name.casefold() in used:
                name=f'{suffix}-{original}'; suffix+=1
            used.add(name.casefold()); archive.writestr(name,content)
    return output.getvalue(),'application/zip','入库原件.zip'


class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args): pass
    def send(self,content,status=200,ctype='application/json; charset=utf-8',filename=None,content_encoding=None,csp=None):
        if not isinstance(content,bytes): content=json.dumps(content,ensure_ascii=False).encode()
        self.send_response(status); self.send_header('Content-Type',ctype); self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff')
        if content_encoding:
            self.send_header('Content-Encoding',content_encoding); self.send_header('Vary','Accept-Encoding')
        self.send_header('Content-Security-Policy',csp or "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'")
        if filename: self.send_header('Content-Disposition',f'attachment; filename="download{Path(filename).suffix}"; filename*=UTF-8\'\'{quote(filename,safe="")}')
        self.send_header('Content-Length',str(len(content))); self.end_headers(); self.wfile.write(content)
    def do_GET(self):
        path=urlparse(self.path).path; query=parse_qs(urlparse(self.path).query)
        if self.headers.get('Host','').split(':')[0] not in ('127.0.0.1','localhost'): return self.send({'error':'Local access only'},403)
        try:
            if path=='/':
                self.send_response(302); self.send_header('Location','/showcase/?warehouse=1'); self.end_headers(); return
            if path=='/showcase':
                self.send_response(302); self.send_header('Location','/showcase/?warehouse=1'); self.end_headers(); return
            if path=='/api/project-arrivals':
                with connect() as c: return self.send(project_arrivals(c))
            if path=='/api/state':
                state=dict(overview(),token=TOKEN,version=APP_VERSION,demo_mode=demo_mode())
                state['admin_password_configured']=admin_password_configured()
                state['local_ai_ready']=local_ai.available()
                return self.send(state)
            if path=='/api/ledger':
                kind=clean(query.get('kind',[''])[0]).upper()
                text=clean(query.get('q',[''])[0])[:100]
                try: limit=int(query.get('limit',[LEDGER_PAGE_SIZE])[0])
                except (TypeError,ValueError): limit=LEDGER_PAGE_SIZE
                try: offset=int(query.get('offset',['0'])[0])
                except (TypeError,ValueError): offset=0
                with connect() as c: return self.send(ledger_page(c,limit,offset,kind,text))
            if path=='/api/items':
                text=clean(query.get('q',[''])[0])[:100]
                package=clean(query.get('package',[''])[0])[:200]
                try: limit=int(query.get('limit',[ITEMS_PAGE_SIZE])[0])
                except (TypeError,ValueError): limit=ITEMS_PAGE_SIZE
                try: offset=int(query.get('offset',['0'])[0])
                except (TypeError,ValueError): offset=0
                with connect() as c: return self.send(items_page(c,limit,offset,text,package))
            if path=='/api/recognitions': return self.send(recognition_list(query))
            if path=='/api/recognition': return self.send(recognition_detail(query['id'][0]))
            if path=='/api/recognition/reopen': return self.send(recognition_reopen(query['id'][0]))
            if path=='/api/stock-detail': return self.send(stock_detail(query.get('id',[''])[0],offset=query.get('offset',['0'])[0]))
            if path=='/api/document': return self.send(doc_detail(query['id'][0]))
            if path=='/api/recognition-source':
                record=recognition_detail(query['id'][0]); rel=record.get('source_path')
                if not rel: raise ValueError('没有原始文件 / Source file not found')
                content,mime,name=source_download(rel,record['original_name'])
                return self.send(content,ctype=mime,filename=name)
            if path=='/api/source':
                d=doc_detail(query['id'][0]); rel=d.get('source_path')
                if not rel: raise ValueError('没有来源文件 / No source file')
                content,mime,name=source_download(rel)
                return self.send(content,ctype=mime,filename=name)
            if path=='/api/backup':
                mem=io.BytesIO()
                with connect() as c:
                    tmp=DATA/('backup-'+secrets.token_hex(6)+'.db')
                    dest=sqlite3.connect(tmp); c.backup(dest); dest.close()
                    try:
                        with zipfile.ZipFile(mem,'w',zipfile.ZIP_DEFLATED) as z:
                            z.write(tmp,'materials.db')
                            for f in (DATA/'sources').glob('*'): z.write(f,'sources/'+f.name)
                    finally: tmp.unlink(missing_ok=True)
                return self.send(mem.getvalue(),ctype='application/zip',filename='materials-backup.zip')
            if path=='/api/export':
                state=overview(); out=io.StringIO(); w=csv.writer(out); w.writerow(['Contract / 合同','Code / 编码','Name / 名称','Specification / 规格','Unit / 单位','Attribute / 属性','Batch / 批次','Box / 箱号','Warehouse / 仓库','Bin / 库位','State / 状态','Quantity / 数量','Received total / 累计入库','Available / 当前可用','Pending / 待检','Quarantine / 隔离','Spares / 备品'])
                for r in state['stock']:
                    vals=[r[k] for k in ('package','code','name','spec','unit','attribute','batch','box','warehouse','bin','state','qty','received_total','available_qty','pending_qty','quarantine_qty','spare_qty')]
                    w.writerow(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v for v in vals])
                return self.send(('\ufeff'+out.getvalue()).encode(),ctype='text/csv; charset=utf-8',filename='inventory.csv')
            if path=='/api/template':
                return self.send('\ufeff物资编码,物资名称,规格型号,单位,数量,批次,箱号,仓库,库位,库存状态,标段\r\nDEMO-001,模拟支架配件,M12,件,100,B01,C01,主仓库,A区-01,可用,A1\r\n'.encode(),ctype='text/csv; charset=utf-8',filename='sample.csv')
            if path=='/api/supplier-template':
                f=ROOT/'static'/'到货单与装箱单模板.xlsx'
                if not f.exists(): raise FileNotFoundError('模板尚未生成 / Template not found')
                return self.send(f.read_bytes(),ctype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='supplier-material-template.xlsx')
            if path=='/api/inbound-demo-data' and demo_mode():
                return self.send((ROOT/'static'/'演示数据-虚构勿入账.xlsx').read_bytes(),ctype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='DEMO-ONLY.xlsx')
            if path in ('/warehouse','/warehouse/'):
                asset=(ROOT/'static'/'index.html').read_bytes()
                return self.send(asset,ctype='text/html; charset=utf-8')
            if path.startswith('/showcase/'):
                showcase=(ROOT/'static'/'showcase').resolve()
                relative=unquote(path[len('/showcase/'):])
                asset_path=(showcase/relative).resolve()
                if asset_path!=showcase and showcase not in asset_path.parents:
                    return self.send({'error':'Not found'},404)
                if asset_path.is_dir(): asset_path=asset_path/'index.html'
                if not asset_path.is_file(): return self.send({'error':'Not found'},404)
                asset=asset_path.read_bytes()
                mime=mimetypes.guess_type(asset_path.name)[0] or 'application/octet-stream'
                if mime.startswith(('text/','application/javascript')): mime += '; charset=utf-8'
                csp="default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'" if asset_path==showcase/'index.html' else None
                if asset_path.suffix in ('.js','.css') and 'gzip' in self.headers.get('Accept-Encoding','').lower():
                    return self.send(gzip.compress(asset,compresslevel=6,mtime=0),ctype=mime,content_encoding='gzip')
                return self.send(asset,ctype=mime,csp=csp)
            names={'/app.js':'app.js','/style.css':'style.css','/tokens.css':'tokens.css','/univer-preview.js':'univer-preview.js','/univer-preview.css':'univer-preview.css','/univer-LICENSES.txt':'univer-LICENSES.txt','/Apache-2.0.txt':'Apache-2.0.txt'}
            if path in names:
                mime={'/':'text/html; charset=utf-8','/app.js':'text/javascript; charset=utf-8','/style.css':'text/css; charset=utf-8','/tokens.css':'text/css; charset=utf-8','/univer-preview.js':'text/javascript; charset=utf-8','/univer-preview.css':'text/css; charset=utf-8','/univer-LICENSES.txt':'text/plain; charset=utf-8','/Apache-2.0.txt':'text/plain; charset=utf-8'}[path]
                asset=(ROOT/'static'/names[path]).read_bytes()
                if path.endswith(('.js','.css')) and 'gzip' in self.headers.get('Accept-Encoding','').lower():
                    return self.send(gzip.compress(asset,compresslevel=6,mtime=0),ctype=mime,content_encoding='gzip')
                return self.send(asset,ctype=mime)
            self.send({'error':'Not found'},404)
        except (ValueError,KeyError,FileNotFoundError) as e: self.send({'error':str(e)},400)
    def do_POST(self):
        if self.headers.get('Host','').split(':')[0] not in ('127.0.0.1','localhost') or self.headers.get('X-Local-Token')!=TOKEN: return self.send({'error':'请刷新页面 / Refresh page'},403)
        origin=self.headers.get('Origin')
        if origin and urlparse(origin).netloc!=self.headers.get('Host'): return self.send({'error':'Origin rejected'},403)
        try:
            length=int(self.headers.get('Content-Length','0'))
            if not 0<length<18*1024*1024: raise ValueError('请求过大 / Request too large')
            p=json.loads(self.rfile.read(length))
            handlers={'/api/preview':preview_post,'/api/intake-ai':intake_ai_post,'/api/import':import_post,'/api/recognition/cancel':cancel_recognition,'/api/recognition/delete':delete_recognition,'/api/recognition/delete-batch':delete_recognitions,'/api/post':stock_post,'/api/reverse':reverse,'/api/delete':admin_delete,'/api/delete-batch':admin_delete_batch,'/api/admin-password':admin_password_set,'/api/change':change_design,'/api/batch':batch_update}
            fn=handlers.get(urlparse(self.path).path)
            if not fn: return self.send({'error':'Not found'},404)
            self.send(fn(p))
        except sqlite3.IntegrityError: self.send({'error':'重复单号、文件或记录冲突，未重复入账 / Duplicate reference, file or conflicting record; not posted'},409)
        except (ValueError,KeyError,TypeError) as e: self.send({'error':str(e)},400)
        except Exception as e:
            print(type(e).__name__,str(e),file=sys.stderr)
            self.send({'error':'处理失败，未完成操作 / Operation failed'},500)

if __name__=='__main__':
    init(); port=int(os.environ.get('MATERIALS_PORT','8765'))
    print(f'Materials local app: http://127.0.0.1:{port}',flush=True)
    ThreadingHTTPServer(('127.0.0.1',port),Handler).serve_forever()
